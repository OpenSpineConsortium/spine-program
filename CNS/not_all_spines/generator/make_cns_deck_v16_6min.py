"""paper/cns2026/make_cns_deck_v16_6min.py -- v16 of the 6-minute CNS deck: v14 with slide 5 replaced (Greg, 2026-10-31).

v15's slide 5 (make_cns_deck_v15_6min.py) with the panels of prep_v16_benzel_profile.py: Panjabi's 12 cadavers drawn
beside Benzel and the gold tier, and each gold tier line drawn only where bordercall's engine read the measure by its
own definition, which leaves out the neck's widths. Built from v14 as v15 was, so the hand-shortened notes of slides 2
and 3 stay, and checked against v14's FORBIDDEN words.

    python paper/cns2026/prep_v16_benzel_profile.py
    python paper/cns2026/make_cns_deck_v16_6min.py
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from prep_v16_benzel_profile import EXTRA, PANELS, panjabi_table, stats  # noqa: E402

DECK = HERE.parent
SOURCE = DECK / "CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx"
TARGET = DECK / "CNS2026_Schwing_NotAllSpines_public_v16_6min.pptx"
FIG = HERE / "figs_v16_6min"
P = json.load(open(FIG / "profile.json", encoding="utf-8"))

# the v14 build's list (make_cns_deck_v14_6min.FORBIDDEN), copied so this runs without v14's imports
FORBIDDEN = ("SchwingNet", "C-arm", "C arm", "carm", "sweep", "proprietary", "fluoro", "render", "DRR", "registration-free",
             "cone-beam", "local shape", "count from the sacrum", "abstain", "abstention", "prove" + "nance", "coupling", "coupled",
             "decoder", "encoder", "stop rule", "chain", "alphabet", "ordinal", "ResNet", "YOLO", "device card",
             "device model", "embedding", "t-SNE", "occlusion", "TotalSegmentator", "BUU", "anchor", "backbone", "neural")
WPM, PAUSE = 140, 3                         # the v14 pace


def level(key, name):
    return P["measures"][key][name]


def drawn(name, key):
    """The levels the panel draws the gold tier's line at for KEY, by the figure's own rule."""
    ticks, place = PANELS[name]["axis"]()
    return [lv for lv in sorted(place, key=place.get) if lv not in EXTRA and stats(P, key, lv)]


def gap(name, key, measure, levels):
    """The largest and the mean distance, in mm, between the gold tier's mean and Panjabi's at LEVELS."""
    table = panjabi_table()
    gaps = [abs(stats(P, key, lv)[0] - table[(measure, lv)][0]) for lv in levels if (measure, lv) in table]
    return max(gaps), sum(gaps) / len(gaps)


def words():
    n = P["spines"]["all"]
    t13, l6 = stats(P, "body_w_mid", "T13"), stats(P, "body_w_mid", "L6")
    canal = [lv for lv in drawn("canal", "interpedicular_max") if lv.startswith(("T", "L"))]
    cw_max, cw_mean = gap("canal", "interpedicular_max", "canal_w", canal)
    lumbar = [lv for lv in drawn("canal", "canal_ap") if lv.startswith("L") and lv != "L6"]
    cd_max, _ = gap("canal", "canal_ap", "canal_ap", lumbar)
    body_depth = [lv for lv in drawn("body", "body_d_mid") if lv != "L6"]
    bd_max, bd_mean = gap("body", "body_d_mid", "endplate_depth", body_depth)
    title = f"The same plots, over {n:,} full spines"
    caption = (f"Black: Benzel, digitised (permission pending). Red: Panjabi, 12 cadavers. Blue: {n:,} whole-spine "
               "CTs, \u00b1 1 SD.")
    alts = (("Benzel's vertebral body width and depth by level in black, Panjabi's cadaveric means with one SD in red, "
             f"and the mean and one SD of {n:,} whole-spine CTs in blue from T1 down for width and at every level for "
             "depth, T13 and L6 as separate points."),
            ("Benzel's spinal canal width and depth by level in black, Panjabi's cadaveric means with one SD in red, "
             f"and the mean and one SD of {n:,} whole-spine CTs in blue from T2 to L4 for width and at every level for "
             "depth, T13 and L6 as separate points."))
    refs = ("Benzel EC. Biomechanics of Spine Stabilization, 3rd ed. Thieme; 2015, ch. 1, Figs 1.1 and 1.9, digitised.   "
            "Panjabi MM, et al. Spine 1991;16:861-869 and 888-901; Spine 1992;17:299-306.   Eisenstein S. J Bone "
            "Joint Surg Br 1977;59:173-180.   Gold tiers 3 and 4: public CT, C2 to the sacrum in view.")
    script = (f"Same two plots. Black is Benzel. Red is Panjabi's twelve cadavers; in the neck Benzel's canal "
              f"figure is Panjabi's, number for number. Blue is {n:,} CT scans that show the whole spine, one scan a "
              f"person, one standard deviation shaded. From the upper thorax down, the canal's width sits within "
              f"{cw_max:.0f} millimeters of Panjabi at every level, and the body's depth within {bd_max:.0f}. The "
              f"textbook drew the average and no spread. And it has no place for a thirteenth thoracic or a sixth "
              f"lumbar vertebra: here they are, {t13[2]} and {l6[2]} spines.")
    asked = ("If asked about the gaps: the neck's widths are not drawn. In the neck the measuring code finds the "
             "foramen transversarium closed in under one vertebra in ten, so it cuts the body at the canal's walls "
             "and its width becomes the canal's plus two millimetres; and it finds both pedicles in about one in a "
             "hundred, so the canal's width becomes the hole's width through its middle. That is our code, not the "
             "labels: every label source and every slice thickness gives the same numbers. Where the pedicles are "
             "found the width is the greatest distance between them (Eisenstein 1977, Panjabi's definition), "
             f"{cw_mean:.1f} mm from Panjabi on average. The canal's depth reads 1 to 5 mm under Panjabi, most in the neck; "
             f"in the lumbar spine it matches the six CT series of the literature pool (15 to 17 mm), and Panjabi's "
             f"cadavers read up to {cd_max:.0f} mm deeper. The T1 and L5 body widths take in the transverse processes' roots. T13 "
             f"rests on {t13[2]} spines and is shown for its place, not its number.")
    return title, caption, alts, refs, script, asked


def set_text(frame, text):
    """The frame's first run takes the text, so the slide's own type and colour stay; the other runs go."""
    paragraphs = frame.paragraphs
    first = paragraphs[0]
    runs = first.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)


def paragraph_text(paragraph, text):
    """A Word paragraph's first run takes the text and keeps its style; the other runs go."""
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def build():
    title, caption, alts, refs, script, asked = words()
    prs = Presentation(str(SOURCE))
    model = [sh for sh in prs.slides[3].shapes if sh.shape_type == 13]
    s = prs.slides[4]
    shapes = list(s.shapes)
    assert shapes[0].text_frame.text.startswith("End-plate width"), shapes[0].text_frame.text
    assert len(model) == 2, len(model)
    old = shapes[1]
    old._element.getparent().remove(old._element)
    for where, spec, alt in zip(model, PANELS.values(), alts):
        pic = s.shapes.add_picture(str(FIG / spec["out"]), where.left, where.top, where.width, where.height)
        pic._element.nvPicPr.cNvPr.set("descr", alt)
    set_text(shapes[0].text_frame, title)
    set_text(shapes[2].text_frame, caption)
    set_text(shapes[3].text_frame, refs)
    notes = s.notes_slide.notes_text_frame
    set_text(notes, script)
    notes.add_paragraph().text = asked
    prs.save(str(TARGET))
    check(TARGET)
    instructions(script, asked)
    for line in (TARGET.name, title, f"{caption} ({len(caption.split())} words)",
                 f"{script} ({len(script.split())} words)", asked):
        print(line)


def instructions(script, asked):
    """v14's instructions with slide 5's row, the times after it, the deck's name and the figure table brought up to v16."""
    from docx import Document
    doc = Document(str(DECK / "CNS2026_NotAllSpines_v14_6min_INSTRUCTIONS.docx"))
    order = doc.tables[0]
    times = [int(row.cells[2].text.split()[0]) for row in order.rows[1:]]
    times[4] = round(len(script.split()) / WPM * 60) + PAUSE
    clock = 0
    for row, sec in zip(order.rows[1:], times):
        clock += sec
        row.cells[2].paragraphs[0].runs[0].text = f"{sec}\n({clock // 60}:{clock % 60:02d})"
    row = order.rows[5]
    row.cells[1].paragraphs[0].runs[0].text = "Benzel and Panjabi, with the gold tiers"
    row.cells[3].paragraphs[0].runs[0].text = script
    head = doc.paragraphs[1]
    text = head.text
    text = re.sub(r"public_v14_6min", "public_v16_6min", text)
    text = re.sub(r"Script \d+ words: about \d+:\d+", f"Script {sum(len(r.cells[3].text.split()) for r in order.rows[1:])} words: "
                  f"about {clock // 60}:{clock % 60:02d}", text)
    text = re.sub(r"so \d+ s in hand", f"so {360 - clock} s in hand", text)
    paragraph_text(head, text)
    bullet = doc.paragraphs[7]
    paragraph_text(bullet, bullet.text + " Slide 5 redraws Figs 1.1 and 1.9 on their own axes, with their lines digitised "
                   "from the book's figures and the seventh tick of Fig. 1.1 labelled 1 (L1), under your gold tier lines, "
                   "so the same Thieme permission covers it. " + asked)
    figures = doc.tables[1].rows[4]
    figures.cells[1].paragraphs[0].runs[0].text = "Benzel with the gold tiers, canal rows, angles"
    figures.cells[2].paragraphs[0].runs[0].text = ("Slide 5: Benzel Figs 1.1 and 1.9 redrawn, his lines digitised (Thieme "
                                                    "permission as slide 4), with your gold tier 3 and 4 measurements. "
                                                    "Canal and angles are in arXiv v1")
    doc.save(str(DECK / "CNS2026_NotAllSpines_v16_6min_INSTRUCTIONS.docx"))


def check(path):
    with zipfile.ZipFile(path) as z:
        text = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    text = re.sub(r"<[^>]+>", " ", text)
    for word in FORBIDDEN:
        assert word.lower() not in text.lower(), word


if __name__ == "__main__":
    build()
