"""paper/cns2026/make_cns_deck_v22_6min.py -- v22 of the 6-minute CNS deck (Greg, 2026-11-01).

Greg's asks of the morning after v21: the pan film had missed the thoracic spine, so it is made again from frames that
show it (grid/pan_render.py, render_v22_pan.py); the four-point naming plot is replaced by the smooth build-up figure
that reaches 100 and 99 per cent (prep_v22_bodies.py); the canal plot from the dataset paper moves up from the backup
into the talk; a new slide turns all eight anomalous spines of the consortium's gallery in 3D, its gallery page linked
(render_v22_anomalies.py); and the end slide shows the consortium's website itself (prep_v22_site.py).

The order is v21's with two slides in: the anomalies after the labels that slide a level, and the canal level by level
after the spinopelvic angles, just before Benzel's one line per level. The scripts of slides 5, 7 and 15 (the angles,
Benzel's own plot and the masks) are cut down to pay for the two new slides; the hand-shortened scripts of slides 2
and 3 stay as Greg left them. Built from v14 as v21 was, through v21's own steps.

    python paper/cns2026/prep_v21_benzel_profile.py
    python paper/cns2026/render_v21_fingerprint.py
    python paper/cns2026/prep_v22_bodies.py
    python paper/cns2026/render_v22_anomalies.py
    python paper/cns2026/prep_v22_site.py
    python paper/cns2026/pan_link.py figs_v22_6min/pan9 --det det_low_L --floor 0.1 --flip det_low_R
    python paper/cns2026/render_v22_pan.py figs_v22_6min/pan9
    python paper/cns2026/make_cns_deck_v22_6min.py
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from make_cns_deck_v20_6min import clone, morph, set_text  # noqa: E402
from make_cns_deck_v21_6min import (DECK, SITE, SOURCE, TIERS, TITLE, VIDEO_BOX, WPM, check, data, draw_state,  # noqa: E402
                                    link_site, name_pictures, notes, order, references, set_paragraph, spinopelvic,
                                    text_slide, video, words)
from prep_v21_benzel_profile import BOXES, STATES  # noqa: E402
from prep_v22_bodies import BOX as PLOT_BOX  # noqa: E402

TARGET = DECK / "CNS2026_Schwing_NotAllSpines_public_v22_6min.pptx"
FIG = HERE / "figs_v22_6min"
V21 = HERE / "figs_v21_6min"
RUN = FIG / "pan9"                                             # the grid run the pan film was drawn from
GALLERY = SITE + "/data/gallery/"
FILM_HEIGHT = 6230000 - 1420000
FILM_BOX = ((12192000 + 1750000 - FILM_HEIGHT * 16 // 9) // 2, 1420000, FILM_HEIGHT * 16 // 9, FILM_HEIGHT)
SCRIPTS = {
    "angles": ("The same masks give the spinopelvic angles, in 756 of 802 patients: pelvic incidence 52.1 degrees, "
               "against 52.1 and 53.4 in the two published supine CT series, and sacral slope and pelvic tilt as "
               "close. The dashed line is standing film, a different posture."),
    "canal": ("And the canal, level by level from T11 to L5: depth holds near 17 millimeters, width opens from 20 to "
              "30, and every median sits inside Panjabi's interval, now with a spread of its own."),
    "benzel": ("This is the textbook: Benzel, chapter 1. Body width and depth on the left, the canal on the right, one "
               "line per level, from cadaveric series of ten to thirty specimens, and no spread drawn."),
    "masks": ("So what do masks buy you? Masked CT goes in. Vertebrae come out: detected, predicted and classified, "
              "including the sixth lumbar the segmenter could not name. Each dot is one vertebra, coloured by its true "
              "level, and the levels sort themselves."),
}


def more_words():
    t = words()
    linked = json.load(open(RUN / "linked.json", encoding="utf-8"))
    counted = [b["name"] for b in linked["bodies"] if b["how"] == "count"]
    t["anomalies"] = dict(
        title="Not all spines are created equal",
        refs=("Turn each one yourself: openspineconsortium.com/data/gallery.   Gold: what sets each spine apart.   "
              "Surfaces from CTSpinoPelvic1K's labels (TCIA CT COLONOGRAPHY)."),
        script=("Here are eight of them: a thirteenth rib-bearing level, a process meeting the ala, a part-fused "
                "transitional vertebra, a fusion surgery made, and twelfth ribs short, one-sided or missing. Each "
                "turns in our online gallery."),
        asked=("If asked: the surfaces are the gallery's own, decimated from the released segmentations; the titles "
               "say what is visible and name no level, because a transitional vertebra cannot be named from a field "
               "that does not reach C2."))
    t["bodies"] = dict(
        title="More bodies in view, surer names",
        caption=("Typical synthetic spines from the measured population, 2,000 of them, as the measures are added up to "
                 "all 32: per cent of fields named exactly, named at 99 per cent certainty, and right when so named; "
                 "1, 3 and 9 bodies in view."),
        refs="Synthetic validation, bordercall build-up study, records 2026-10-28; unpublished.   " + TIERS,
        script=("The more of a spine you see, the surer the name. With three bodies and all 32 measures, 98 per cent "
                "are named exactly; with nine, every one, at 99 per cent certainty, and every one right."),
        asked=("If asked: these are typical columns drawn from the gold tier's own measures and correlations; the "
               "rule reads every body in view at once. Of fields it names at 0.99, 99.6 to 100 per cent are right at "
               "every point drawn."))
    shaped = [b["name"] for b in linked["bodies"] if b["how"] == "shape"]
    named = (f"all {len(counted)} named bodies, L5 to {counted[-1]}, are counted up from the L5 it named itself"
             if not shaped else f"{len(counted)} bodies are counted up from the L5 it named itself and "
             f"{len(shaped)} named by its own per-frame reads")
    t["pan"]["asked"] = (
        "If asked: the frames are made from a public CT, not a patient exposure: its shoulder bones softened, the "
        "shoulders cut away and the thoracic bodies weighted up, as a swimmer's view would show them. Boxes are the "
        "detector's own, read at a lower floor over each frame and its mirror image, and kept only on the spine's "
        f"column; {named}, every one as the CT labels it.")
    t["pan"]["script"] = ("Last, a lateral pan up from the sacrum, prone, as in the room. Each body is found as it "
                          "comes into view and named counting up from the sacrum, through the thoracic spine to the "
                          "neck, and the three-dimensional map fills in as it goes.")
    return t


def anomalies_slide(prs, model, t):
    a = t["anomalies"]
    slide = text_slide(prs, model, a["title"], None, a["refs"])
    refs = [sh for sh in slide.shapes if sh.has_text_frame][-1]
    run = refs.text_frame.paragraphs[0].runs[0]
    head, link, tail = "Turn each one yourself: ", "openspineconsortium.com/data/gallery", run.text.split("gallery", 1)[1]
    run.text = head
    linked = refs.text_frame.paragraphs[0].add_run()
    linked.text = link
    linked.hyperlink.address = GALLERY
    rest = refs.text_frame.paragraphs[0].add_run()
    rest.text = tail
    for r in (linked, rest):
        r.font.size, r.font.name = run.font.size, run.font.name
        if run.font.color and run.font.color.type is not None:
            r.font.color.rgb = run.font.color.rgb
    seconds = json.load(open(FIG / "anomalies.json", encoding="utf-8"))["seconds"]
    movie = video(slide, FIG / "anomalies.mp4", FIG / "anomalies_first.png", FILM_BOX, seconds, "anomalies_video")
    movie.click_action.hyperlink.address = GALLERY
    notes(slide, a)
    return slide


def site_window(end):
    """The website in its browser window where the mark stood, the window itself (not its shadow) exactly the mark's
    width; the links and the questions box move down under it."""
    from PIL import Image
    from prep_v22_site import SHADOW
    logo = [sh for sh in end.shapes if sh.shape_type == 13][0]
    left, top, width = logo.left, logo.top - 60000, logo.width
    logo._element.getparent().remove(logo._element)
    im = Image.open(FIG / "site_window.png")
    w = int(width * im.width / (im.width - 2 * SHADOW))
    h = int(w * im.height / im.width)
    margin = w * SHADOW // im.width
    pic = end.shapes.add_picture(str(FIG / "site_window.png"), left - margin, top - margin, w, h)
    pic.click_action.hyperlink.address = SITE
    pic._element.nvPicPr.cNvPr.set("descr", "The Open Spine Consortium's website, openspineconsortium.com, its home page.")
    bottom = top - margin + h - margin
    links = next(sh for sh in end.shapes if sh.has_text_frame and "Weights and code" in sh.text_frame.text)
    links.top = bottom + 90000
    box = next(sh for sh in end.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Questions"))
    box.top = links.top + links.height + 120000


def script(slide, key):
    set_paragraph(slide.notes_slide.notes_text_frame.paragraphs[0], SCRIPTS[key])


def build():
    t = more_words()
    store = (*data(), references())
    prs = Presentation(str(SOURCE))
    s = list(prs.slides)
    title, cost, labels, benzel, first, canal_rows, angles, masks, end = s
    name_pictures(benzel)
    old = first.shapes[1]
    assert first.shapes[0].text_frame.text.startswith("End-plate width"), first.shapes[0].text_frame.text
    old._element.getparent().remove(old._element)
    morphs = [first]
    for _ in STATES[1:]:
        morphs.append(clone(prs, morphs[-1]))
    fp, bd, pn = t["fingerprint"], t["bodies"], t["pan"]
    film = text_slide(prs, first, fp["title"], fp["caption"], fp["refs"])      # cloned while first holds text alone
    bodies = text_slide(prs, first, bd["title"], bd["caption"], bd["refs"])
    pan = text_slide(prs, first, pn["title"], None, pn["refs"])
    gallery = anomalies_slide(prs, first, t)
    for slide, state in zip(morphs, STATES):
        shapes = [sh for sh in slide.shapes if sh.has_text_frame]
        set_text(shapes[0].text_frame, TITLE)
        set_text(shapes[1].text_frame, t[state]["caption"])
        set_text(shapes[2].text_frame, t[state]["refs"])
        draw_state(slide, state, store)
        notes(slide, t[state])
        morph(slide)
    spinopelvic(angles)
    script(angles, "angles")
    script(benzel, "benzel")
    script(masks, "masks")
    script(canal_rows, "canal")
    video(film, V21 / "fingerprint.mp4", V21 / "fingerprint_first.png", VIDEO_BOX, 18.8, "!!fingerprint_video")
    notes(film, fp)
    morph(film)
    last = bodies.shapes.add_picture(str(V21 / "fingerprint_last.png"), *VIDEO_BOX)
    last.name = "!!fingerprint_video"
    plot = bodies.shapes.add_picture(str(FIG / "bodies_curve.png"), *PLOT_BOX)
    plot.name = "!!bodies_plot"
    plot._element.nvPicPr.cNvPr.set("descr", "Three panels against the measures added, 1 to 32, for 1, 3 and 9 bodies in "
                                    "view: per cent named exactly, per cent named at 99 per cent certainty, and per cent "
                                    "right of those so named, from 99 to 100.")
    notes(bodies, bd)
    morph(bodies)
    video(pan, FIG / "pan.mp4", FIG / "pan_first.png", FILM_BOX,
          json.load(open(FIG / "pan.json", encoding="utf-8"))["seconds"], "pan_video")
    notes(pan, pn)
    link_site(end)
    site_window(end)
    order(prs, [title, cost, labels, gallery, angles, canal_rows, benzel, *morphs, film, bodies, masks, pan, end])
    prs.save(str(TARGET))
    check(TARGET)
    rows = timing(prs)
    instructions(rows, TARGET.stat().st_size / 1e6)
    print(TARGET.name, f"{TARGET.stat().st_size / 1e6:.0f} MB")


def timing(prs):
    """Each slide's words at 140 a minute plus 3 s for the change, a film's slide at least its film and 2 s."""
    films = {"!!fingerprint_video": 18.8, "pan_video": json.load(open(FIG / "pan.json", encoding="utf-8"))["seconds"],
             "anomalies_video": json.load(open(FIG / "anomalies.json", encoding="utf-8"))["seconds"]}
    clock, rows = 0.0, []
    for i, slide in enumerate(prs.slides, 1):
        said = slide.notes_slide.notes_text_frame.paragraphs[0].text if slide.has_notes_slide else ""
        sec = len(said.split()) / WPM * 60 + 3
        for shape in slide.shapes:
            if shape.shape_type == 16:
                sec = max(sec, films.get(shape.name, 0) + 2)
        clock += sec
        title = next((sh.text_frame.text for sh in slide.shapes if sh.has_text_frame), "")
        rows.append((i, title, sec, clock, said))
        print(f"{i:2d} {sec:5.0f} s  {int(clock // 60)}:{clock % 60:04.1f}  {title[:60]}")
    return rows


def instructions(rows, megabytes):
    """v21's instructions brought to v22: the running order slide by slide, the check list and the figures."""
    from docx import Document
    doc = Document(str(DECK / "CNS2026_NotAllSpines_v21_6min_INSTRUCTIONS.docx"))
    table = doc.tables[0]
    for row in list(table.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    merged = []
    for r in rows:
        if 8 <= r[0] <= 12:
            if r[0] == 8:
                merged.append(["8-12", TITLE + ": five clicks", 0.0, 0.0, []])
            merged[-1][2] += r[2]
            merged[-1][3] = r[3]
            merged[-1][4].append(r[4])
        else:
            merged.append([str(r[0]), r[1], r[2], r[3], [r[4]]])
    for number, title, sec, clock, said in merged:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = number, title
        cells[2].text = f"{sec:.0f}\n({int(clock // 60)}:{int(clock % 60):02d})"
        cells[3].text = " [click] ".join(x for x in said if x)
    end = rows[-1][3]
    words = sum(len(r[4].split()) for r in rows)
    head = doc.paragraphs[1]
    text = head.text.replace("public_v21_6min", "public_v22_6min")
    text = re.sub(r"\d+ slides \(the last a backup after the end\)", f"{len(rows)} slides, no backup", text)
    text = re.sub(r"Script \d+ words: about \d+:\d+", f"Script {words} words: about {int(end // 60)}:{int(end % 60):02d}",
                  text)
    text = re.sub(r"and the two films at their own length", "and the three films at their own length", text)
    text = re.sub(r"so \d+ s (in hand against 6:00|over 6:00[^.]*)",
                  f"so {360 - end:.0f} s in hand against 6:00" if end <= 360 else
                  f"so {end - 360:.0f} s over 6:00: cut slide 3's last two sentences first", text)
    set_paragraph(head, text)
    for paragraph in doc.paragraphs:
        old = paragraph.text
        new = (old.replace("slides 13 and 14", "slides 15 and 16").replace("Slide 14's", "Slide 16's")
               .replace("Slides 5 to 12 are joined by Morph", "Slides 7 to 14 are joined by Morph")
               .replace("Slides 6 to 10 are one slide in five clicks", "Slides 8 to 12 are one slide in five clicks")
               .replace("The films on 11 and 14", "The films on 4, 13 and 16")
               .replace("with the two films", "with the three films"))
        if new != old:
            set_paragraph(paragraph, new)
    figures = doc.tables[1]
    for row in list(figures.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    for cells in (("4", "Anomalies film", "The consortium gallery's eight anomalous spines, from CTSpinoPelvic1K's "
                   "released labels", "Yours; the gallery page is linked on the slide"),
                  ("6", "Canal level by level", "Your dataset article's Fig. 3, moved up from the backup", "Yours"),
                  ("7-12", "Benzel with the gold tiers", "Benzel Figs 1.1 and 1.9 themselves as the base, your gold "
                   "tier 3 and 4 lines and Panjabi's, Tan's and Alonge's published means over them", "Thieme, as before"),
                  ("13", "Fingerprint film", "Your gold tier decades in three dimensions; four spines built from one "
                   "ENHANCE.PET CT", "Yours; ENHANCE.PET is CC BY and CC BY-NC, cited on the slide"),
                  ("14", "Naming as measures are added", "bordercall's build-up study, records of 28 October", "Yours"),
                  ("15", "Closing picture", "Your outcomes on public CT (CTSpinoPelvic1K, VerSe), cited", "Yours; filed"),
                  ("16", "Pan film", "One ENHANCE.PET CT, FLUX.2 klein-9B (FLUX non-commercial licence), your detector",
                   "Non-commercial talk: within both licences"),
                  ("17", "End slide", "The consortium's website, its home page as Edge shows it", "Yours")):
        row = figures.add_row().cells
        for cell, text in zip(row, cells):
            cell.text = text
    doc.save(str(DECK / "CNS2026_NotAllSpines_v22_6min_INSTRUCTIONS.docx"))


if __name__ == "__main__":
    build()
