"""paper/cns2026/make_cns_v14_6min_instructions.py -- the two-page instructions for the 6-minute deck: running order with
seconds, the spoken script per slide (the same text as the speaker notes, imported, never retyped), what to check
before the session, and which figures need permission. Black tables, no colour. Same disclosure check as the deck.

    python paper/cns2026/make_cns_v14_6min_instructions.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from make_cns_deck_v14_6min import FORBIDDEN, ROOT, SCRIPT, TIMES

OUT = ROOT.parent / "CNS2026_NotAllSpines_v14_6min_INSTRUCTIONS.docx"
FONT = "Arial"
SIZE = 9


def run(p, text, bold=False, size=SIZE):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    return r


def para(doc_or_cell, text="", bold=False, size=SIZE, after=2, before=0, align=None):
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    if text:
        run(p, text, bold, size)
    return p


def heading(doc, text):
    p = para(doc, text, bold=True, size=11, after=2, before=6)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    run(p, lead, True)
    run(p, text)
    return p


def borders(table):
    """Plain black rules, the house style for tables."""
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def fill(table, rows, widths, head=True):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            c = table.cell(i, j)
            c.width = Inches(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run(p, str(text), bold=head and i == 0)
    for j, w in enumerate(widths):
        table.columns[j].width = Inches(w)


def mmss(s):
    return f"{s // 60}:{s % 60:02d}"


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.65)
    sec.top_margin, sec.bottom_margin = Inches(0.55), Inches(0.5)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(SIZE)
    cp = doc.core_properties
    cp.author = "Gregory J. Schwing"; cp.last_modified_by = "Gregory J. Schwing"
    cp.title = "CNS 2026, 6-minute deck: instructions"; cp.comments = ""

    para(doc, "Not All Spines Are Created Equal: instructions for the 6-minute deck", bold=True, size=13, after=1)
    total = sum(TIMES)
    para(doc, "CNS 2026, Mini Podium 4, Spine Session 1 (MP4SP1), Monday 2 November 2026, 5:21–6:00 PM, Podium 4, Ballroom ABC. "
              f"Deck: CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx, 9 slides. Script {sum(len(t.split()) for _, t in SCRIPT)} words: "
              f"about {mmss(total)} at 140 words a minute with 3 s per slide change, so 45 s in hand against 6:00. "
              "The same script is in the speaker notes; open Presenter View.", after=4)

    heading(doc, "Running order and script")
    rows = [("#", "Slide", "Sec (clock)", "Say")]
    clock = 0
    for i, ((name, text), sec_) in enumerate(zip(SCRIPT, TIMES), 1):
        clock += sec_
        rows.append((i, name, f"{sec_}\n({mmss(clock)})", text))
    t = doc.add_table(rows=len(rows), cols=4)
    fill(t, rows, [0.25, 0.95, 0.55, 5.45])

    heading(doc, "Check before the session")
    bullet(doc, "Slide 8 is the one slide not yet public. ",
           "Decide it first. It shows outcomes and no method, but the TTO has not cleared Part B. If it is not cleared, delete slide 8: "
           "the deck then runs " + mmss(total - TIMES[7]) + " and ends on the spinopelvic slide, then the end.")
    bullet(doc, "Questions on slide 8. ", "Say what it does, never how. A ready line: \"That work is unpublished, so I will leave it there today.\"")
    bullet(doc, "Slide 3. ", "The lettered picture you pasted is not on disk. This one is cut from the three panels of NeurIPS Figure 1 and re-ordered to match it "
                             "(A CT, B segmenter, C truth), with the letters added. Compare the two by eye.")
    bullet(doc, "Slide 4. ", "Fig. 1.1 prints \"12\" where L1 belongs on its axis (Fig. 1.9 has it right); that is the book, not the deck. "
                             "You asked for one Benzel figure and Fig. 1.1 alone filled under half the slide, so Fig. 1.9 sits beside it and sets up slide 6. Drop it if you prefer one permission.")
    bullet(doc, "ACCME. ", "No disclosure slide, as in v13: CNS shows its own list. If a reported financial relationship has changed, email education@cns.org now. "
                           "CNS reviews the slides. Say \"the segmenter\", not its brand name; the slides use generic names.")
    bullet(doc, "Practice. ", "Time one run aloud in Presenter View. If it runs past 5:45, cut the last two sentences of slide 3 first.")
    bullet(doc, "File. ", "Under 10 MB, so it goes by email; keep a copy on a USB stick too.")

    heading(doc, "Figures and permission")
    rows = [("#", "Figure", "Source", "Permission"),
            ("2", "Cost board", "Your composition of published numbers; no publisher figure copied", "None needed; citations on the slide"),
            ("3", "CT panels; 3-D view", "Your NeurIPS 2026 Figure 1; your dataset article Fig. 2(d), in arXiv v1", "Yours. Check the journal's terms if Medical Physics accepts the article"),
            ("4", "Benzel Figs 1.1 and 1.9", "Thieme, 2015, chapter 1. All rights reserved; the chapter PDF is licensed through WSU", "NEEDED. Marked \"permission pending\" on the slide. Ask Thieme"),
            ("5–7", "Violins, canal rows, angles", "Your data, drawn with the article's code. Canal and angles are in arXiv v1; the end-plate violin is new as a figure, its data are public", "Yours"),
            ("8", "Closing picture", "Your unpublished outcomes on public CT (CTSpinoPelvic1K, VerSe), cited on the slide", "No publisher. TTO decision, above")]
    t = doc.add_table(rows=len(rows), cols=4)
    fill(t, rows, [0.42, 1.25, 3.25, 2.28])

    # the same disclosure check as the deck
    blob = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for tb in doc.tables for r in tb.rows for c in r.cells)
    hits = [w for w in FORBIDDEN if w.lower() in blob.lower()]
    assert not hits, hits
    doc.save(str(OUT))
    print("wrote", OUT, "words", len(blob.split()))


if __name__ == "__main__":
    build()
