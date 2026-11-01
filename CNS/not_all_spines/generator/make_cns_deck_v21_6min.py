"""paper/cns2026/make_cns_deck_v21_6min.py -- v21 of the 6-minute CNS deck (Greg, 2026-10-31, night).

Greg's asks of the night, in his order: the populations on one slide through Morph (purple, the gold tier reweighted
to Panjabi's sex and age; men and women apart; Tan's Asian series; Alonge's African series; then every decade of each
sex, blue and red, deeper the farther from his cadavers' age); Benzel's two plots at exactly slide 4's size and place,
his vertebra drawing never under a line; then the plot turned into depth, four ideal spines lit gold one by one with
their lines, the lines leaving as snakes and settling as a fingerprint (one video, render_v21_fingerprint.py); the
generator's naming on the 32 measures as more bodies come into view, beside that fingerprint; the masks slide; the
pan from the sacrum with the detector working and the three-dimensional map filling in (render_v21_pan.py), patent
pending on the slide, nothing written in the video; and the questions and acknowledgements with the consortium's link.

Slide 5 is v14's slide 4 with its two pictures named so Morph holds them; slides 6 to 10 lay their lines over the same
two scans in the same boxes (prep_v21_benzel_profile.py has the calibration) and carry one title. The spinopelvic
slide moves ahead of Benzel and the canal rows go after the end as a backup, so the new sequence runs unbroken.
Built from v14, as v15 to v20 were, so the hand-shortened notes of slides 2 and 3 stay.

    python paper/cns2026/prep_v21_benzel_profile.py
    python paper/cns2026/render_v21_fingerprint.py
    python paper/cns2026/prep_v21_bodies.py
    python paper/cns2026/render_v21_pan.py
    python paper/cns2026/make_cns_deck_v21_6min.py
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from make_cns_deck_v20_6min import FORBIDDEN, clone, diamond, morph, plain, polyline, set_text  # noqa: E402
from prep_v21_benzel_profile import (BOXES, CIRCLE_LINE_PT, CIRCLE_PT, EXTRA_INK, REF_PT, REFERENCE, STATES,  # noqa: E402
                                     Panel, data, extras, lines, phantom, points, references)

DECK = HERE.parent
SOURCE = DECK / "CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx"
TARGET = DECK / "CNS2026_Schwing_NotAllSpines_public_v21_6min.pptx"
FIG = HERE / "figs_v21_6min"
OLD = HERE / "figs_v14_6min"
V20 = HERE / "figs_v20_6min"
WPM = 140
SITE = "https://openspineconsortium.com"
VIDEO_BOX = (BOXES["body"][0], BOXES["body"][1], BOXES["canal"][0] + BOXES["canal"][2] - BOXES["body"][0],
             BOXES["body"][3])
TITLE = "Benzel's plots, redrawn from our whole-spine CTs"
BENZEL = "Benzel EC. Biomechanics of Spine Stabilization, 3rd ed. Thieme; 2015, ch. 1, Figs 1.1 and 1.9, digitised.   "
TIERS = "Gold tiers 3 and 4: public CT, C2 to the sacrum in view."
PANJABI = "Panjabi MM, et al. Spine 1991;16:861-869 and 888-901; Spine 1992;17:299-306.   "
AUTOPLAY = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
            '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" '
            'nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            '<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/>'
            '<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>'
            '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            '</p:seq><p:video><p:cMediaNode vol="0"><p:cTn id="7" fill="hold" display="0"><p:stCondLst>'
            '<p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            '</p:cMediaNode></p:video></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')


def words():
    rw, bands, _ = data()
    n = rw["people"]
    tan = rw["targets"]["tan"]["who"]
    pj = rw["targets"]
    gap = pj["panjabi_men"]["measures"]["plate_w_l"]["L3"]["mean"] - pj["panjabi_women"]["measures"]["plate_w_l"]["L3"]["mean"]
    least, most = min(bands["counts"].values()), max(bands["counts"].values())
    return {
        "panjabi": dict(
            caption=(f"Black: Benzel, digitised (permission pending). Gray: Panjabi, 12 cadavers. Purple: {n:,} adults' "
                     "whole-spine CTs reweighted to his sex and age. Open circles: T13 and L6, every spine."),
            refs=BENZEL + PANJABI + TIERS,
            script=(f"Benzel's plots again, with our data on them. Gray is Panjabi's twelve cadavers. Purple is {n:,} "
                    "adults' whole-spine CTs, reweighted to his sex and age: the canal's width sits within a millimeter "
                    "of his. And the textbook has no place for a thirteenth thoracic or a sixth lumbar: here they are."),
            asked=("If asked: each of Panjabi's twelve specimens is matched to our adults of its sex within five years "
                   "of its age, and the line is the mean of the twelve. T13 and L6 are every spine's means.")),
        "sexes": dict(
            caption="Blue: our men, matched by age to Panjabi's 8 men. Red: our women, matched to his 4 women.",
            refs=BENZEL + PANJABI + TIERS,
            script=f"Divide it by sex: men run {gap:.0f} millimeters wider than women at L3, and his line sits between.",
            asked="If asked: Panjabi gives no means by sex; these are the same matching over his men and his women alone."),
        "tan": dict(
            caption=(f"Gray: Tan, 10 Chinese Singaporean cadavers, 56 to 77 years. Green: our {tan['people']} adults "
                     "inside their age, height and weight."),
            refs="Tan SH, Teo EC, Chua HC. Eur Spine J 2004;13:137-146.   " + BENZEL + TIERS,
            script="An Asian series, Tan's cadavers, beside our adults of their age and build: smaller bodies.",
            asked=(f"If asked: Tan reports no sex; our {tan['people']} adults in their ranges are {tan['men']} men. Their "
                   "canal was measured on dry bone with a digitiser, a third shallower than Panjabi's.")),
        "alonge": dict(
            caption=("Gray: Alonge, 120 Nigerian adults' CTs, L1 to L5; width at the widest pedicle cut. Orange: our "
                     "adults reweighted to their sex and age."),
            refs="Alonge OJ, Adeolu AA, Omololu AB, Atalabi OM. Afr J Med Med Sci 2021;50:69-76.   " + BENZEL + TIERS,
            script="An African series, Alonge's Nigerian adults on CT, against ours reweighted to them.",
            asked=("If asked: no African American series in the literature I hold measures these chords; the Hamann-Todd "
                   "skeletons give pedicles only. Alonge's width is interpedicular, which is why it climbs to L5.")),
        "agesex": dict(
            caption=("Every decade apart, 20s to 80s. Blue, men; red, women; the deeper the colour, the farther the "
                     f"decade's mean age from Panjabi's cadavers' 46. {least} to {most} adults a line; canal depth not "
                     "shown by decade."),
            refs=BENZEL + PANJABI + TIERS,
            script=("Now every decade of each sex. The farther from his cadavers' age, the deeper the colour, and "
                    "the further the line strays from his."),
            asked=("If asked: the decades are unweighted means of every adult of that sex and decade, ten or more a "
                   "level. Canal depth is left out by decade because on the CT it reads deeper with age, 16 to 23 mm "
                   "at L3, the opposite of what age does: the reading leaks in older bone, and is being fixed.")),
        "fingerprint": dict(
            title="Each spine, its own fingerprint",
            caption=("Each decade's width and depth as one curve; four spines built from one CT, widened and deepened "
                     "to their decade's means, each lit with its line; then every line becomes a ridge."),
            refs="Spines: one ENHANCE.PET 1.6k whole-body CT, its vertebrae scaled level by level.   " + TIERS,
            script=("Turn the plot into depth and the decades come apart. Each line is a spine of its own. Every "
                    "spine has its own fingerprint."),
            asked=("If asked: the four spines are one CT's vertebrae, each scaled about its own centre by its decade's "
                   "mean width and depth over every decade's mean; the fingerprint is a drawing of the lines, not data.")),
        "bodies": dict(
            title="More bodies in view, surer names",
            caption=("Synthetic spines from the measured population, all 32 measures, 200 fields a point: per cent of "
                     "fields whose middle body is named exactly, by the bodies in view."),
            refs="Synthetic validation, bordercall records 2026-10-28; unpublished.   " + TIERS,
            script=("Because each spine has its own fingerprint, the more of it you see, the surer the name: one body, "
                    "85 per cent; three, 95; five, 98, with six lumbar bodies or a thirteenth thoracic as well."),
            asked=("If asked: the columns are drawn from the gold tier's own measures and their correlations; the "
                   "rule reads every body in view at once.")),
        "pan": dict(
            title="Panning up from the sacrum",
            refs=("Patent pending.   X-ray frames made from one public whole-body CT the model never saw "
                  "(ENHANCE.PET 1.6k)."),
            script=("Last, a lateral pan up from the sacrum, prone, as in the room. Each body is found as it comes "
                    "into view and named counting up from the sacrum, and the three-dimensional map fills in as it goes."),
            asked=("If asked: the frames are made from a public CT, not a patient exposure, and say so in words if "
                   "anyone asks. Boxes are the detector's own; names below the shoulders are counted from the L5 it "
                   "named itself, and above the shoulders are its own per-frame reads; T1 and T2, under the shoulders, "
                   "are left unnamed.")),
    }


def name_pictures(slide):
    for shape, name in zip([sh for sh in slide.shapes if sh.shape_type == 13], BOXES):
        shape.name = f"!!{name}_base"


def circle(slide, name, centre):
    size = Pt(CIRCLE_PT)
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(centre[0] - size / 2), int(centre[1] - size / 2), int(size), int(size))
    plain(shape)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor.from_string(EXTRA_INK[1:])
    shape.line.width = Pt(CIRCLE_LINE_PT)
    return shape


def draw_state(slide, state, store):
    rw, bands, profile, refs = store
    panels = [Panel(name) for name in BOXES]
    for panel in panels:
        pic = slide.shapes.add_picture(str(panel.image), *panel.box)
        pic.name = f"!!{panel.name}_base"
    for panel in panels:
        for chord, pts in points(panel, state, refs).items():
            runs = panel.clip([p for _, p in pts]) if len(pts) > 1 else []
            for i, run in enumerate(runs):                       # a published value off Benzel's axis is cut there
                polyline(slide, f"!!{panel.name}_r{chord}" + (f"_{i + 1}" if i else ""), [panel.emu(p) for p in run],
                         REFERENCE, REF_PT, MSO_LINE.ROUND_DOT)
            for level, p in pts:
                if panel.inside(p):
                    diamond(slide, f"!!{panel.name}_r{chord}_{level}", panel.emu(p), REFERENCE)
        for name, p in extras(panel, profile):
            circle(slide, f"!!{name}", panel.emu(p))
        for name, colour, pts, width, dashed in lines(panel, state, rw, bands):
            for i, run in enumerate(panel.clip(pts)):
                polyline(slide, f"!!{name}" + (f"_{i + 1}" if i else ""), [panel.emu(p) for p in run], colour, width,
                         MSO_LINE.DASH if dashed else None)
    for panel in panels:
        path, (x0, y0, x1, y1) = phantom(panel)
        a, b = panel.emu((x0, y0)), panel.emu((x1, y1))
        pic = slide.shapes.add_picture(str(path), a[0], a[1], b[0] - a[0], b[1] - a[1])
        pic.name = f"!!{panel.name}_phantom"
        pic._element.nvPicPr.cNvPr.set("descr", "Benzel's vertebra drawing, laid over the lines so none hides it.")


def autoplay(slide, movie, seconds):
    """Play the movie as the slide opens, with no click, and hold its last frame."""
    root = slide._element
    for old in root.findall(qn("p:timing")):
        root.remove(old)
    timing = etree.fromstring(AUTOPLAY.replace("{spid}", str(movie.shape_id)).replace("{ms}", str(int(seconds * 1000))))
    ext = root.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(timing)
    else:
        root.append(timing)


def video(slide, path, poster, box, seconds, name):
    movie = slide.shapes.add_movie(str(path), *box, poster_frame_image=str(poster), mime_type="video/mp4")
    movie.name = name
    autoplay(slide, movie, seconds)
    return movie


def text_slide(prs, model, title, caption, refs):
    slide = clone(prs, model)
    shapes = [sh for sh in slide.shapes if sh.has_text_frame]
    set_text(shapes[0].text_frame, title)
    if caption is None:
        shapes[1]._element.getparent().remove(shapes[1]._element)
    else:
        set_text(shapes[1].text_frame, caption)
    set_text(shapes[2].text_frame, refs)
    return slide


def notes(slide, t):
    frame = slide.notes_slide.notes_text_frame
    frame.text = t["script"]
    frame.add_paragraph().text = t["asked"]


def spinopelvic(slide):
    pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
    old = pictures[0]
    pic = slide.shapes.add_picture(str(V20 / "spinopelvic_benzel.png"), old.left, old.top, old.width, old.height)
    pic._element.nvPicPr.cNvPr.set("descr", "Histograms of pelvic incidence, sacral slope and pelvic tilt over the gold "
                                   "tier CTs, with the published supine CT and standing film means.")
    old._element.getparent().remove(old._element)
    caption = [sh for sh in slide.shapes if sh.has_text_frame][1]
    set_text(caption.text_frame, json.load(open(V20 / "spinopelvic.json", encoding="utf-8"))["caption"])


def link_site(slide):
    """The consortium's mark takes the audience to the site, and the site heads the links under it."""
    logo = [sh for sh in slide.shapes if sh.shape_type == 13][0]
    logo.click_action.hyperlink.address = SITE
    links = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Weights and code"))
    first = links.text_frame.paragraphs[0]
    para = first._p.addprevious(etree.fromstring(etree.tostring(first._p)))
    para = links.text_frame.paragraphs[0]
    for run in para.runs[1:]:
        run._r.getparent().remove(run._r)
    run = para.runs[0]
    run.text = "openspineconsortium.com"
    run.hyperlink.address = SITE


def order(prs, slides):
    ids = prs.slides._sldIdLst
    by_id = {e.get("id"): e for e in ids}
    wanted = [by_id[str(s.slide_id)] for s in slides]       # read before the list is emptied
    for e in list(ids):
        ids.remove(e)
    for e in wanted:
        ids.append(e)


def build():
    t = words()
    store = (*data(), references())
    prs = Presentation(str(SOURCE))
    s = list(prs.slides)
    title, cost, labels, benzel, first, canal_rows, angles, masks, end = s
    name_pictures(benzel)
    old = first.shapes[1]
    assert first.shapes[0].text_frame.text.startswith("End-plate width"), first.shapes[0].text_frame.text
    old._element.getparent().remove(old._element)
    morphs = [first]
    for _ in STATES[1:]:
        morphs.append(clone(prs, morphs[-1]))
    fp, bd, pn = t["fingerprint"], t["bodies"], t["pan"]
    film = text_slide(prs, first, fp["title"], fp["caption"], fp["refs"])      # cloned while first holds text alone
    bodies = text_slide(prs, first, bd["title"], bd["caption"], bd["refs"])
    pan = text_slide(prs, first, pn["title"], None, pn["refs"])
    for slide, state in zip(morphs, STATES):
        shapes = [sh for sh in slide.shapes if sh.has_text_frame]
        set_text(shapes[0].text_frame, TITLE)
        set_text(shapes[1].text_frame, t[state]["caption"])
        set_text(shapes[2].text_frame, t[state]["refs"])
        draw_state(slide, state, store)
        notes(slide, t[state])
        morph(slide)
    spinopelvic(angles)
    video(film, FIG / "fingerprint.mp4", FIG / "fingerprint_first.png", VIDEO_BOX, 18.8, "!!fingerprint_video")
    notes(film, fp)
    morph(film)
    last = bodies.shapes.add_picture(str(FIG / "fingerprint_last.png"), *VIDEO_BOX)
    last.name = "!!fingerprint_video"
    plot = bodies.shapes.add_picture(str(FIG / "bodies_benzel.png"), *BOXES["canal"])
    plot.name = "!!bodies_plot"
    plot._element.nvPicPr.cNvPr.set("descr", "Per cent of synthetic fields whose middle body is named exactly against the "
                                    "bodies in view, 1 to 9, for typical spines, six lumbar bodies and a thirteenth "
                                    "thoracic, all 32 measures.")
    notes(bodies, bd)
    morph(bodies)
    height = 6230000 - 1420000
    width = height * 16 // 9
    video(pan, FIG / "pan.mp4", FIG / "pan_first.png", ((12192000 + 1750000 - width) // 2, 1420000, width, height),
          json.load(open(FIG / "pan.json", encoding="utf-8"))["seconds"], "pan_video")
    notes(pan, pn)
    link_site(end)
    order(prs, [title, cost, labels, angles, benzel, *morphs, film, bodies, masks, pan, end, canal_rows])
    prs.save(str(TARGET))
    check(TARGET)
    instructions(timing(prs), TARGET.stat().st_size / 1e6)
    print(TARGET.name, f"{TARGET.stat().st_size / 1e6:.0f} MB")


def check(path):
    with zipfile.ZipFile(path) as z:
        text = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    text = re.sub(r"<[^>]+>", " ", text)
    for word in FORBIDDEN:
        assert word.lower() not in text.lower(), word


def timing(prs):
    """Each slide's words at 140 a minute plus 3 s for the change, a video slide at least its film's length; the
    running clock; printed and returned for the instructions."""
    clock, rows = 0.0, []
    films = {"!!fingerprint_video": 18.8, "pan_video": json.load(open(FIG / "pan.json", encoding="utf-8"))["seconds"]}
    for i, slide in enumerate(prs.slides, 1):
        said = slide.notes_slide.notes_text_frame.paragraphs[0].text if slide.has_notes_slide else ""
        sec = len(said.split()) / WPM * 60 + 3
        for shape in slide.shapes:
            if shape.shape_type == 16 or getattr(shape, "media_type", None) is not None:
                sec = max(sec, films.get(shape.name, 0) + 2)
        clock += sec
        title = next((sh.text_frame.text for sh in slide.shapes if sh.has_text_frame), "")
        rows.append((i, title, sec, clock, said))
        print(f"{i:2d} {sec:5.0f} s  {int(clock // 60)}:{clock % 60:04.1f}  {title[:60]}")
    return rows


def instructions(rows, megabytes):
    """v20's instructions brought to v21: the running order rebuilt slide by slide, the check list and the figures."""
    from docx import Document
    doc = Document(str(DECK / "CNS2026_NotAllSpines_v20_6min_INSTRUCTIONS.docx"))
    order = doc.tables[0]
    for row in list(order.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    talk = [r for r in rows if r[0] <= 15]
    merged = []
    for r in talk:
        if 6 <= r[0] <= 10:
            if r[0] == 6:
                merged.append(["6-10", TITLE + ": five clicks", 0.0, 0.0, []])
            merged[-1][2] += r[2]
            merged[-1][3] = r[3]
            merged[-1][4].append(r[4])
        else:
            merged.append([str(r[0]), r[1], r[2], r[3], [r[4]]])
    for number, title, sec, clock, said in merged:
        cells = order.add_row().cells
        cells[0].text, cells[1].text = number, title
        cells[2].text = f"{sec:.0f}\n({int(clock // 60)}:{int(clock % 60):02d})"
        cells[3].text = " [click] ".join(x for x in said if x)
    end = talk[-1][3]
    words = sum(len(r[4].split()) for r in talk)
    head = doc.paragraphs[1]
    text = head.text.replace("public_v20_6min", "public_v21_6min")
    text = re.sub(r"\d+ slides", "16 slides (the last a backup after the end)", text)
    text = re.sub(r"Script \d+ words: about \d+:\d+ at 140 words a minute with 3 s per slide change",
                  f"Script {words} words: about {int(end // 60)}:{int(end % 60):02d} at 140 words a minute with 3 s per "
                  "slide change and the two films at their own length", text)
    text = re.sub(r"so \d+ s in hand against 6:00", f"so {360 - end:.0f} s in hand against 6:00" if end <= 360 else
                  f"so {end - 360:.0f} s over 6:00: cut slide 4's script to its first sentence first", text)
    set_paragraph(head, text)
    checks = [
        ("Disclosure. The provisional is filed (your word of 31 October), so slides 13 and 14 are shown. Slide 14's "
         "foot says patent pending; nothing is written in either film."),
        ("Questions on slides 13 and 14. Say what it does; the method is in the filing. A ready line: \"That is in a "
         "pending application, so I will leave the how there today.\""),
        None, None, None,
        ("Morph and the films. Slides 5 to 12 are joined by Morph (PowerPoint 2019 or Microsoft 365; an older one "
         "fades). Slides 6 to 10 are one slide in five clicks. The films on 11 and 14 start by themselves when the "
         "slide opens and hold their last frame; click on when each ends. Check both in the speaker-ready room."),
        (f"File. About {megabytes:.0f} MB with the two films, too big to email: bring it on a USB stick and keep a "
         "copy in OneDrive."),
    ]
    for paragraph, text in zip(doc.paragraphs[4:11], checks):
        if text:
            set_paragraph(paragraph, text)
    honest = doc.paragraphs[10].insert_paragraph_before(
        "Slide 14's frames are made from a public whole-body CT, not a patient exposure, made to look like an X-ray "
        "by FLUX.2 klein-9B as released; its foot says so in a few words. If anyone asks, say it plainly.")
    honest.style = doc.paragraphs[10].style
    figures = doc.tables[1]
    for row in list(figures.rows)[1:]:
        if row.cells[0].text.startswith(("5", "11")):
            row._tr.getparent().remove(row._tr)
    for cells in (("5-10", "Benzel with the gold tiers", "Benzel Figs 1.1 and 1.9 themselves as the base, your gold "
                   "tier 3 and 4 lines and Panjabi's, Tan's and Alonge's published means over them", "Thieme, as slide 4"),
                  ("11", "Fingerprint film", "Your gold tier decades in three dimensions; four spines built from one "
                   "ENHANCE.PET CT", "Yours; ENHANCE.PET is CC BY and CC BY-NC, cited on the slide"),
                  ("12", "Bodies in view", "bordercall's synthetic validation, records of 28 October", "Yours"),
                  ("13", "Closing picture", "Your outcomes on public CT (CTSpinoPelvic1K, VerSe), cited", "Yours; filed"),
                  ("14", "Pan film", "One ENHANCE.PET CT, FLUX.2 klein-9B (FLUX non-commercial licence), your detector",
                   "Non-commercial talk: within both licences")):
        row = figures.add_row().cells
        for cell, text in zip(row, cells):
            cell.text = text
    doc.save(str(DECK / "CNS2026_NotAllSpines_v21_6min_INSTRUCTIONS.docx"))


def set_paragraph(paragraph, text):
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


if __name__ == "__main__":
    build()
