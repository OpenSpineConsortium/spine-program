"""paper/cns2026/make_cns_deck_v14_6min.py -- the CNS 2026 talk cut to six minutes (Greg, 2026-10-30): mostly pictures,
legible words, same CNS template and style as public v13, in his order: the money slide, the three-panel CT with a 3-D
view, Benzel, his per-level plots, the spinopelvic angles, then a picture that shows but does not tell, and the end.

Disclosure rule (binding): slide 8 shows outcomes as pictures and says what, never how. The build fails if a word from
FORBIDDEN appears in any slide text, note, alt text or document property. Every number on a slide or in the script is
asserted below against the file it comes from: the NeurIPS camera-ready PDF, the MDPA manuscript, arXiv v1 of the
dataset article, level_atlas.csv, or the numbers prep_v14_6min_figs.py measured from the CSVs.

Stays out of git: this folder is excluded in .git/info/exclude of the public repository.

    python paper/cns2026/prep_v14_6min_figs.py
    python paper/cns2026/make_cns_deck_v14_6min.py
    python paper/cns2026/make_cns_v14_6min_instructions.py
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import zipfile
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from make_cns_deck_v2 import FIGS2, fit
from make_cns_deck_v4 import (BLUE, GREY, L, NAVY, R, REFS_TOP, TEMPLATE, TOP, WHITE, add_text, content_slide, navy_bar,
                              picture, refs)
from make_cns_deck_v6 import notes

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
FIG = HERE / "figs_v14_6min"
N = json.load(open(FIG / "numbers.json", encoding="utf-8"))

TITLE = "Not All Spines Are Created Equal: A Single-Shot CT Segmentation Model For Preoperative Detection Of Lumbar And Thoracic End Anomalies"
SESSION = "Mini Podium 4, Spine Session 1 (MP4SP1)  ·  Monday 2 November 2026, 5:21–6:00 PM  ·  Podium 4, Ballroom ABC"
PAPER = "Not All Spines Are Created Equal, NeurIPS 2026, Evaluations and Datasets track (accepted)"
REPO = "https://github.com/OpenSpineConsortium/spinopelvic-seg"
REF_NEURIPS = "Schwing GJ, et al. Not All Spines Are Created Equal: How CT Segmenters Fail on Lumbosacral Transitional Vertebrae. NeurIPS 2026 (accepted)."
REF_DATA = "CTSpinoPelvic1K, doi.org/10.5281/zenodo.22139642."
ARXIV = "Dataset article: arXiv 2609.22760"

CAP = 14                  # one caption line under a figure, slide type
FIG_BOTTOM = 5250000      # full-width figures end here; the caption line under them ends by 5.65 M EMU
TEXT_LIMIT = 5650000
EMU_W = R - L

# Words that may not appear anywhere in the file. "render" stays banned: the 3-D view is called a 3-D view.
FORBIDDEN = ("SchwingNet", "C-arm", "C arm", "carm", "sweep", "proprietary", "fluoro", "render", "DRR", "registration-free",
             "cone-beam", "local shape", "count from the sacrum", "abstain", "abstention", "prove" + "nance", "coupling", "coupled",
             "decoder", "encoder", "stop rule", "chain", "alphabet", "ordinal", "ResNet", "YOLO", "device card",
             "device model", "embedding", "t-SNE", "occlusion", "TotalSegmentator", "BUU", "anchor", "backbone", "neural")
# slide 8 may say what, never how: its own text and notes also stay away from these
FORBIDDEN_8 = ("train", "network", "model", "algorithm", "pipeline", "feature", "crop", "projection", "radiograph", "film",
               "x-ray", "xray", "exposure", "patent")

# seconds and the spoken script, one entry per slide: shared with the instructions document
SCRIPT = [
    ("Title", "I'm Greg Schwing, a general surgery resident at Wayne State and the Detroit Medical Center. "
              "Not all spines are created equal, and the ones that are not are expensive."),
    ("Money slide", "Start with the money. US malpractice payments for surgical never events reached 1.3 billion dollars, and "
                    "one in four of those paid claims is wrong site. A closed wrong-site claim averages 136 thousand dollars, "
                    "and spine is the commonest specialty in that series. In litigated wrong-disc-level cases, 81 percent were "
                    "paid. Nearly one spine surgeon in five has been sued for wrong level or side. The grey tiles are my own estimate: "
                    "420 to 1,900 wrong-level spine operations a year in the United States. And prevention is not free: a gold "
                    "fiducial costs close to two thousand dollars a patient."),
    ("One CT, two readings, 3-D", "Here is one reason. One patient. A is the CT. C is the truth: six lumbar bodies, the lowest in "
                                  "red. B is the most widely used open-source CT segmenter, which has no sixth lumbar class. It "
                                  "fits six bodies into five labels, so below L2 every label is one level off: well formed, "
                                  "confident, and no error signal. On normal anatomy its L5 Dice is 0.94. On lumbarization it is "
                                  "0.29. Across 1,152 CTs it emitted L6 zero times. On the right is a six-lumbar spine in 3-D from "
                                  "CTSpinoPelvic1K, the public dataset I built to measure this: 802 patients, spine, pelvis, ribs "
                                  "and femora in one frame, 33 transitional spines graded by Castellvi. The single-shot model I "
                                  "released merges L5 and L6 in training, which repairs the junction and the sacrum. It does not "
                                  "yet tell four lumbar bodies from five."),
    ("Benzel", "Those five labels assume a textbook spine. This is the textbook: Benzel, chapter 1. Vertebral body width and depth "
               "on the left, the canal on the right, one line per level. These lines come from cadaveric series of ten to thirty "
               "specimens, and no spread is drawn. They cannot tell you whether the patient in front of you is unusual."),
    ("End-plate width by level", "Here is the same idea from 802 CT scans: upper end-plate width, T11 to L4, one violin per level. "
                                 "In red is Panjabi's cadaveric interval, twelve specimens, his standard errors converted back to a "
                                 "standard deviation. From L1 to L4 my medians sit within half a millimeter of his means, which is "
                                 "the check. What the textbook never showed is the width of each violin. At L4 the middle 90 percent "
                                 "of patients runs from 39 to 56 millimeters. L5 is withheld on purpose: the transverse processes "
                                 "fall in front of the cut that isolates the body, so that number would be wrong."),
    ("Canal by level", "The canal, T11 to L5, this time with L5. Depth is a midsagittal chord at the pedicle, not a maximum "
                       "extent, because that is what the normative series measure. It holds near 17 millimeters at every level. "
                       "Width opens from 20 to 30. Every median sits inside Panjabi's interval, and now each level has its own spread."),
    ("Spinopelvic angles", "The same masks give the spinopelvic angles, in the 756 of 802 patients that pass the geometric gates. "
                           "Pelvic incidence 52.1 degrees, against 52.1 and 53.4 in the two published supine CT series. Sacral slope "
                           "36.6, against 36.5 and 34.1. Pelvic tilt 15.5, against 15.6 and 19.2. The dashed line is standing film, a "
                           "different posture. I call these provisional until the sacral plate comes from a substructure label, but "
                           "the pelvis in these masks measures what published CT measures."),
    ("Masked CT in, vertebrae out", "So what do masks buy you? This. Masked CT goes in. Vertebrae come out: detected, predicted, "
                                    "classified, including the sixth lumbar vertebra the segmenter could not name. Each dot on the "
                                    "right is one vertebra, coloured by its true level, and the levels sort themselves. I am showing "
                                    "you what it does, not how. This part is unpublished work in progress."),
    ("The end", "That is the end. The CT side is public: the data on Zenodo, the weights and code on GitHub. Thank you to the students "
                "and radiology residents who built it, and thank you for listening."),
]
WPM = 140                                   # a measured conference pace
PAUSE = 3                                   # seconds to change slide and let it land


def seconds(text: str) -> int:
    return round(len(text.split()) / WPM * 60) + PAUSE


TIMES = [seconds(t) for _, t in SCRIPT]
TIMES[0] = max(TIMES[0], 12)


def verify_numbers():
    """Every load-bearing number, against the file it comes from."""
    cam = " ".join(" ".join(p.get_text() for p in fitz.open(ROOT / "CAMERA_READY_NEURIPS" / "src" / "neurips_2026.pdf")).split())
    for s in ("0.938", "0.291", "0/1,152"):
        assert s in cam, s
    assert "every label below L2 one level off" in cam
    main = " ".join((ROOT / "paper" / "mpda" / "main.tex").read_text(encoding="utf-8").split())
    for s in (r"Pelvic incidence & 52.1 $\pm$ 11.1\si{\degree} & 52.1 & 53.4",
              r"Sacral slope & 36.6 $\pm$ 10.6\si{\degree} & 36.5 & 34.1",
              r"Pelvic tilt & 15.5 $\pm$ 7.9\si{\degree} & 15.6 & 19.2",
              "over the 756 of 802 records", "cadaveric series of ten to thirty specimens", "are the provisional part of this release",
              "33 cases typed I--IV"):
        assert s in main, s
    ax = " ".join(" ".join(p.get_text() for p in fitz.open(FIG / "arxiv_2609.22760v1.pdf")).split())
    for s in ("Morphometry by level, T11–L5", "Four phenotypes at the lumbar borders", "six lumbar bodies, Castellvi IV",
              "Pelvic incidence 52.1 ± 11.1", "Sacral slope 36.6 ± 10.6", "Pelvic tilt 15.5 ± 7.9"):
        assert s.replace("±", "±") in ax or s.split(" ±")[0] in ax, s
    for key, want in (("pelvic_incidence_deg", [756, 52.1, 11.1]), ("sacral_slope_deg", [756, 36.6, 10.6]),
                      ("pelvic_tilt_deg", [756, 15.5, 7.9])):
        assert N["spinopelvic"][key] == want, key
    atlas = {(r["measure"], r["level"]): r for r in csv.DictReader(open(ROOT / "morphometrics" / "level_atlas.csv"))}
    assert round(float(atlas[("endplate", "L4")]["p5"])) == 39 and round(float(atlas[("endplate", "L4")]["p95"])) == 56
    assert N["violin"]["gap"] <= 0.5 + 1e-9 and (N["violin"]["n_min"], N["violin"]["n_max"]) == (723, 768)
    assert N["canal"]["depth_median"] == [16.4, 17.8] and N["canal"]["width_median"] == [20.3, 30.0]
    assert all(N["canal"]["inside_panjabi_interval"].values()) and len(N["canal"]["inside_panjabi_interval"]) == 14
    board = (HERE / "prep_v14_figs.py").read_text(encoding="utf-8")
    for s in ("$1.3 billion", "one in\\nfour paid never-event claims", "$136,453", "spine the commonest (22 %)", "81 % paid", "1 in 5",
              "420–1,900 a year", "$1,750–1,950"):
        assert s in board, s


def alt(shape, text):
    shape._element.nvPicPr.cNvPr.set("descr", text)
    return shape


def caption(s, y, text, size=CAP, color=GREY):
    add_text(s, L, y, EMU_W, TEXT_LIMIT - y, [text], size=size, color=color, gap=0)


def full_figure(prs, title, img, cap_text, alt_text, refs_text):
    s = content_slide(prs, title)
    w, h = fit(img, EMU_W, FIG_BOTTOM - TOP)
    alt(picture(s, img, L + (EMU_W - w) // 2, TOP, w, h), alt_text)
    caption(s, TOP + h + 50000, cap_text)
    refs(s, refs_text)
    return s


def build(out: Path):
    verify_numbers()
    prs = Presentation(str(TEMPLATE))
    sld = prs.slides._sldIdLst
    RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for e in list(sld):
        sld.remove(e)
        prs.part.drop_rel(e.get(RID))

    # ---- 1 title (as v13: sole author; everyone else under acknowledgments)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    t = ph[0]
    t.left, t.top, t.width, t.height = Emu(5185558), Emu(1050000), Emu(6600000), Emu(2200000)
    t.text_frame.word_wrap = True
    t.text_frame.text = TITLE
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = WHITE
    sub = ph[1]
    sub.left, sub.top, sub.width, sub.height = Emu(5185558), Emu(3350000), Emu(6600000), Emu(900000)
    sub.text_frame.word_wrap = True
    sub.text_frame.text = "Gregory J. Schwing, MD, PhD"
    p = sub.text_frame.add_paragraph()
    p.text = "Department of General Surgery, Detroit Medical Center and Wayne State University, Detroit, MI"
    for i, p in enumerate(sub.text_frame.paragraphs):
        for r in p.runs:
            r.font.size = Pt(20 if i == 0 else 13); r.font.color.rgb = WHITE
    body = ph[10]
    body.left, body.top, body.width, body.height = Emu(5185558), Emu(4400000), Emu(6600000), Emu(1800000)
    body.text_frame.word_wrap = True
    body.text_frame.text = "CNS 2026 Annual Meeting, Washington, DC"
    for line in (SESSION, "Results from: " + PAPER):
        p = body.text_frame.add_paragraph()
        p.text = line
    for p in body.text_frame.paragraphs:
        p.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(12); r.font.color.rgb = WHITE
    notes(s, [SCRIPT[0][1]])

    # ---- 2 the money slide: the cost board of v8-v13, the same image
    s = content_slide(prs, "What a wrong site costs, and what avoiding it costs")
    img = FIGS2 / "board_wrongsite_cost.png"
    w, h = fit(img, EMU_W, REFS_TOP - 100000 - TOP)
    alt(picture(s, img, L + (EMU_W - w) // 2, TOP, w, h), "Published costs of wrong-site surgery and of the checks that prevent it.")
    refs(s, "Mehtsun WT, et al. Surgery 2013;153:465.   Tan et al. Jt Comm J Qual Patient Saf 2023;49:265.   Bettis T, et al. Patient Saf Surg 2026.   Goodkin R, Laska LL. Surg Neurol 2004;61:323.   Mody MG, et al. Spine 2008;33:194.   Groff MW, et al. World Neurosurg 2013;79:585.   Machin JT, et al. Eur Spine J 2018;27:2693.   Harrison WD, et al. Ann R Coll Surg Engl 2015.   AJNR Am J Neuroradiol 2023;44:618.   Childers CP, Maggard-Gibbons M. JAMA Surg 2018;153:e176233.")
    notes(s, [SCRIPT[1][1],
              "If asked: every number is from a verified primary source except the four marked derived: the wrong-site share of the NPDB total (one quarter of paid never-event claims are wrong site), the yearly cost of the fiducial series (179 patients over 38.4 months at $34–37 per operating-room minute, 2014 dollars), and the US wrong-level operation and indemnity estimates (the two published per-procedure rates applied to 1.3–1.6 million spine operations a year, with the published litigation and payment fractions).",
              "The $1.3 billion is all surgical never events (retained objects, wrong site, wrong procedure, wrong patient) over 1990–2010; no published wrong-site-only total exists."])

    # ---- 3 this pic, with a 3-D view
    s = content_slide(prs, "One CT, six lumbar bodies, and the labels slide a level")
    a, r3 = FIG / "pic_abc.png", FIG / "render_l6.png"
    gap = 140000
    ha, wa = None, None
    from PIL import Image
    ra, rr = (lambda z: z[0] / z[1])(Image.open(a).size), (lambda z: z[0] / z[1])(Image.open(r3).size)
    h = int((EMU_W - gap) / (ra + rr))
    wa, wr = int(h * ra), int(h * rr)
    alt(picture(s, a, L, TOP, wa, h), "Coronal CT of one patient in three panels: plain CT; the segmenter's labels, shifted one level, with the sacrum and hips; the true labels, six lumbar bodies with the lowest in red.")
    alt(picture(s, r3, L + wa + gap, TOP, wr, h), "Three-dimensional view of a spine with six lumbar bodies, numbered 1 to 6, and the pelvis.")
    y = TOP + h + 50000
    caption(s, y, "A, CT. B, the segmenter's output. C, ground truth, six lumbar bodies. Right, a six-lumbar spine in 3-D (Castellvi IV).")
    add_text(s, L, y + 300000, EMU_W, TEXT_LIMIT - y - 300000,
             ["L5 Dice 0.94 on normal anatomy, 0.29 on lumbarization.  L6 emitted: 0 of 1,152 CTs."], size=20, gap=0, bold=True)
    refs(s, "Wasserthal J, et al. Radiol Artif Intell 2023;5:e230024.   " + REF_NEURIPS + "   " + ARXIV + ", Fig. 2.   " + REF_DATA)
    notes(s, [SCRIPT[2][1]])

    # ---- 4 Benzel, the embedded images themselves
    s = content_slide(prs, "The plot we learned from: one line per level")
    i1, i2 = FIG / "benzel_fig1_1.png", FIG / "benzel_fig1_9.png"
    bh = FIG_BOTTOM - TOP
    w1, h1 = fit(i1, 4900000, bh)
    w2, h2 = fit(i2, 4900000, bh)
    gapb = 450000
    x1 = L + (EMU_W - w1 - w2 - gapb) // 2
    alt(picture(s, i1, x1, TOP, w1, h1), "Textbook plot of vertebral body width and depth against spinal level, one line each, no spread.")
    alt(picture(s, i2, x1 + w1 + gapb, TOP, w2, h2), "Textbook plot of spinal canal width and depth against spinal level, one line each, no spread.")
    caption(s, TOP + max(h1, h2) + 50000, "Left, vertebral body diameter (Fig. 1.1). Right, spinal canal diameter (Fig. 1.9). Benzel 2015; permission pending.")
    refs(s, "Benzel EC. Biomechanics of Spine Stabilization, 3rd ed. Thieme; 2015, chapter 1, Figs 1.1 and 1.9 (ISBN 978-1-60406-924-2). Figure data: Berry et al; Panjabi et al; White and Panjabi; Reynolds et al; McCormack et al.")
    notes(s, [SCRIPT[3][1]])

    # ---- 5 the violins
    s = full_figure(prs, "End-plate width in 802 patients, with the spread", FIG / "violin_endplate.png",
                    "Retrospective supine CT, n = 723–768 per level. Black tick, median. Panjabi SD recovered as SEM × √12.",
                    "Violin plots of upper end-plate width from T11 to L4, with Panjabi's cadaveric interval in red from L1 to L4; L5 withheld.",
                    "Panjabi MM, et al. Spine 1992;17(3):299–306.   " + ARXIV + ", Fig. 3.   " + REF_DATA)
    notes(s, [SCRIPT[4][1]])

    # ---- 6 the canal rows
    s = full_figure(prs, "The canal, T11 to L5, level by level", FIG / "atlas_canal.png",
                    "Median, interquartile range, 5th–95th percentile, n = 663–784 per level. Depth is the midsagittal chord at the pedicle.",
                    "Canal width and depth by level from T11 to L5, each with its spread, over Panjabi's cadaveric means and interval.",
                    "Panjabi MM, et al. Spine 1991;16(8):888–901 and Spine 1992;17(3):299–306.   " + ARXIV + ", Fig. 3.   " + REF_DATA)
    notes(s, [SCRIPT[5][1]])

    # ---- 7 the spinopelvic angles
    s = full_figure(prs, "Spinopelvic angles match published supine CT", FIG / "spinopelvic.png",
                    "756 of 802 records passing the geometric gates, mean ± SD. Veilleux n = 200, Hasegawa n = 24. Band, standing mean ± SD.",
                    "Distributions of pelvic incidence, sacral slope and pelvic tilt over 756 records, with supine CT and standing reference lines.",
                    "Veilleux NJ, et al. J Bone Joint Surg Am 2020;102:e130.   Hasegawa K, et al. BMC Musculoskelet Disord 2018;19:437.   Vialle R, et al. J Bone Joint Surg Am 2005;87:260.   " + ARXIV + ", Table III and Fig. S2.")
    notes(s, [SCRIPT[6][1]])

    # ---- 8 shows, does not tell
    s = full_figure(prs, "Masked CT in, vertebrae out", FIG / "closing.png",
                    "Masked CT in; vertebrae detected, predicted and classified out. Map colour is true level. Work in progress; not a cleared device.",
                    "Left, a sagittal CT with coloured bone masks. Then the vertebrae of the same spine outlined; six vertebrae with their level names; and a map of vertebrae coloured by level, with the sixth lumbar marked.",
                    "CT and labels: " + REF_DATA + "   VerSe: Sekuboyina A, et al. Med Image Anal 2021;73:102166. Public, de-identified data. Unpublished results.")
    notes(s, [SCRIPT[7][1]])

    # ---- 9 the end: acknowledgments and data, as v13, with v13's evidence-status line
    s = content_slide(prs, "The end. Thank you")
    add_text(s, L, TOP, 6300000, TEXT_LIMIT - TOP, [
        "With thanks to Patrick Schwing (LSU Health Sciences Center New Orleans); Miraziz Ismoilov and Nizar Alnabahneh (Radiology) and Loren Schwiebert (Computer Science), Wayne State University; and the annotators [student A-K]",
        "Data: CTSpinoPelvic1K (doi.org/10.5281/zenodo.22139642), built on CTSpine1K, CTPelvic1K and The Cancer Imaging Archive; all de-identified",
        "Paper: " + PAPER,
        "Funding: NIH NIGMS F30GM147909. No commercial support.",
        "Evidence status: retrospective, public de-identified data; small transitional strata; not a cleared device."], size=16, gap=7)
    logo = FIGS2 / "osc_logo.png"
    lw, lh = fit(logo, 3600000, 1100000)
    pic = alt(picture(s, logo, R - 3600000 + (3600000 - lw) // 2, TOP + 60000, lw, lh), "Open Spine Consortium logo.")
    pic.click_action.hyperlink.address = REPO
    tb = add_text(s, R - 3600000, TOP + lh + 200000, 3600000, 800000,
                  ["Weights and code:", "github.com/OpenSpineConsortium/spinopelvic-seg", "huggingface.co/OpenSpineConsortium"],
                  size=11, gap=3, align=PP_ALIGN.CENTER)
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            if r.text.startswith("github.com"):
                r.hyperlink.address = REPO
            if r.text.startswith("huggingface.co"):
                r.hyperlink.address = "https://huggingface.co/OpenSpineConsortium/spinopelvic-seg-checkpoints"
    navy_bar(s, R - 3600000, TOP + lh + 1150000, 3600000, 1000000, ["Questions", "gregory.schwing@med.wayne.edu"], size=18)
    notes(s, [SCRIPT[8][1]])

    # document properties carry no template author and no working title
    cp = prs.core_properties
    cp.author = "Gregory J. Schwing"
    cp.last_modified_by = "Gregory J. Schwing"
    cp.title = "Not All Spines Are Created Equal"
    cp.subject = ""; cp.keywords = ""; cp.comments = ""; cp.category = ""

    # checks on the finished deck
    assert len(prs.slides) == len(SCRIPT) == 9
    assert sum(TIMES) <= 345, sum(TIMES)                # 5:45 at most, so a slow day still lands inside 6:00
    area_of_content = (R - L) * (5700000 - TOP)
    txt = []
    for n, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else ""
        if n > 1:
            assert len(title) <= 60, title
        per = []
        area = 0
        for sh in sl.shapes:
            if sh.has_text_frame:
                per.append(sh.text_frame.text)
                if sh.top >= REFS_TOP:
                    assert sh.top + sh.height <= 6858000, n
                elif sh.is_placeholder is False and n > 2:
                    # text under a figure stops by the limit
                    assert sh.top + sh.height <= TEXT_LIMIT + 1000, (n, sh.text_frame.text[:50], sh.top + sh.height)
            if sh.shape_type == 13:
                d = sh._element.nvPicPr.cNvPr.get("descr", "")
                assert d and not d.lower().endswith((".png", ".jpg")), n
                per.append(d)
                area += sh.width * sh.height
        if 2 <= n <= 8:
            assert area >= 0.5 * area_of_content, (n, area / area_of_content)
        if sl.has_notes_slide:
            per.append(sl.notes_slide.notes_text_frame.text)
        blob = "\n".join(per)
        txt.append(blob)
        if n == 8:
            low = blob.lower()
            hits = [w for w in FORBIDDEN_8 if w in low]
            assert not hits, hits
    blob = "\n".join(txt).lower()
    hits = [w for w in FORBIDDEN if w.lower() in blob]
    assert not hits, hits
    prs.save(str(out))

    # the saved package: properties, and the size limit for email
    z = zipfile.ZipFile(out)
    meta = (z.read("docProps/core.xml") + z.read("docProps/app.xml")).decode("utf-8", "ignore").lower()
    hits = [w for w in FORBIDDEN if w.lower() in meta]
    assert not hits, hits
    assert "nick gbur" not in meta
    size = out.stat().st_size
    assert size < 10_000_000, size
    mins = sum(TIMES)
    print(f"wrote {out}  slides {len(prs.slides)}  {size / 1e6:.1f} MB  spoken {sum(len(t.split()) for _, t in SCRIPT)} words  "
          f"{mins // 60}:{mins % 60:02d}")
    print(" ", TIMES)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT.parent / "CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx"))
    a = ap.parse_args()
    build(Path(a.out))
