"""paper/cns2026/make_cns_deck_v15_6min.py -- v15 of the 6-minute CNS deck: v14 with slide 5 replaced (Greg, 2026-10-31).

Slide 5 was the end-plate violins. Greg asked (issue #247, and again on 2026-10-31) for the Benzel figures of slide 4
again, the same way, with a second line from his gold tier 3 and 4 CTs, T13 and L6 included, with the standard
deviation and the population named in the legend; prep_v15_benzel_profile.py draws the two panels. Every other slide
is v14's as it stands, the shortened notes of slides 2 and 3 included: this opens the v14 file and changes only slide
5's title, pictures, caption, references and notes, the two pictures where slide 4 has its two, then checks the whole
deck against v14's FORBIDDEN words, as the v14 build does, and brings the instructions up to v15.

    python paper/cns2026/prep_v15_benzel_profile.py
    python paper/cns2026/make_cns_deck_v15_6min.py
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from prep_v15_benzel_profile import PANELS  # noqa: E402

DECK = HERE.parent
SOURCE = DECK / "CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx"
TARGET = DECK / "CNS2026_Schwing_NotAllSpines_public_v15_6min.pptx"
FIG = HERE / "figs_v15_6min"
P = json.load(open(FIG / "profile.json", encoding="utf-8"))

# the v14 build's list (make_cns_deck_v14_6min.FORBIDDEN), copied so this runs without v14's imports
FORBIDDEN = ("SchwingNet", "C-arm", "C arm", "carm", "sweep", "proprietary", "fluoro", "render", "DRR", "registration-free",
             "cone-beam", "local shape", "count from the sacrum", "abstain", "abstention", "prove" + "nance", "coupling", "coupled",
             "decoder", "encoder", "stop rule", "chain", "alphabet", "ordinal", "ResNet", "YOLO", "device card",
             "device model", "embedding", "t-SNE", "occlusion", "TotalSegmentator", "BUU", "anchor", "backbone", "neural")
WPM, PAUSE = 140, 3                         # the v14 pace


def level(key, name):
    return P["measures"][key][name]


def outside():
    """Benzel's points that fall outside the gold tier mean +/- 1 SD, by panel, and how many points there are."""
    out, points = [], 0
    for name, spec in PANELS.items():
        ticks, _ = spec["axis"]()
        for key, word, values in ((spec["keys"][0], "width", spec["benzel_width"]),
                                  (spec["keys"][1], "depth", spec["benzel_depth"])):
            for lv, value in zip(ticks, values):
                points += 1
                row = level(key, lv)
                if abs(value - row["mean"]) > row["sd"]:
                    out.append(f"{name} {word} {lv}")
    return out, points


def words():
    n = P["spines"]["all"]
    t13, l6 = level("body_w_mid", "T13"), level("body_w_mid", "L6")
    missed, points = outside()
    body = [m for m in missed if m.startswith("body")]
    body_points = 2 * len(PANELS["body"]["benzel_width"])
    title = f"The same plots, over {n:,} full spines"
    caption = (f"Black: Benzel, digitised (permission pending). Blue: gold tiers 3 and 4, mean ± 1 SD of {n:,} "
               "whole-spine CTs.")
    alts = (("Benzel's vertebral body width and depth by level in black, with the mean and one standard deviation of "
             f"{n:,} whole-spine CTs in blue at every level, and T13 and L6 as separate points."),
            ("Benzel's spinal canal width and depth by level in black, with the mean and one standard deviation of "
             f"{n:,} whole-spine CTs in blue at every level, and T13 and L6 as separate points."))
    refs = ("Benzel EC. Biomechanics of Spine Stabilization, 3rd ed. Thieme; 2015, chapter 1, Figs 1.1 and 1.9 (ISBN "
            "978-1-60406-924-2), lines digitised.   Gold tiers 3 and 4: public CT with C2 to the sacrum in view "
            "(ENHANCE.PET, VerSe and others). Fig. 1.1's seventh tick is L1.")
    script = (f"Here are the same two plots. Benzel's lines in black. In blue, the same four diameters measured on "
              f"{n:,} CT scans that show the whole spine, one scan a person, with one standard deviation shaded. "
              f"In the vertebral body the textbook lines sit inside that band at {body_points - len(body)} of their "
              f"{body_points} points. "
              f"What the book never drew is the spread, and it has no place for a thirteenth thoracic or a sixth lumbar "
              f"vertebra. Here they are, {t13['n']} and {l6['n']} spines, measured like any other level.")
    asked = ("If asked: Benzel falls outside one SD at " + (", ".join(missed) or "no point") + f" ({len(missed)} of "
             f"{points} points). The CT canal reads narrower and shallower in the neck than his series; the engine's "
             "width through the canal's middle is known to read 2 to 3 mm under the papers' in the neck and at L4 and "
             "L5, and the depth gap in the neck is not explained yet. T1 and L5 body widths carry the transverse "
             f"processes beside the body. T13 rests on {t13['n']} spines and is shown for its place, not its number.")
    return title, caption, alts, refs, script, asked


def set_text(frame, text):
    """The frame's first run takes the text, so the slide's own type and colour stay; the other runs go."""
    paragraphs = frame.paragraphs
    first = paragraphs[0]
    runs = first.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)


def paragraph_text(paragraph, text):
    """A Word paragraph's first run takes the text and keeps its style; the other runs go."""
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def build():
    title, caption, alts, refs, script, asked = words()
    prs = Presentation(str(SOURCE))
    model = [sh for sh in prs.slides[3].shapes if sh.shape_type == 13]
    s = prs.slides[4]
    shapes = list(s.shapes)
    assert shapes[0].text_frame.text.startswith("End-plate width"), shapes[0].text_frame.text
    assert len(model) == 2, len(model)
    old = shapes[1]
    old._element.getparent().remove(old._element)
    for where, spec, alt in zip(model, PANELS.values(), alts):
        pic = s.shapes.add_picture(str(FIG / spec["out"]), where.left, where.top, where.width, where.height)
        pic._element.nvPicPr.cNvPr.set("descr", alt)
    set_text(shapes[0].text_frame, title)
    set_text(shapes[2].text_frame, caption)
    set_text(shapes[3].text_frame, refs)
    notes = s.notes_slide.notes_text_frame
    set_text(notes, script)
    notes.add_paragraph().text = asked
    prs.save(str(TARGET))
    check(TARGET)
    instructions(script, asked)
    for line in (TARGET.name, title, f"{caption} ({len(caption.split())} words)",
                 f"{script} ({len(script.split())} words)", asked):
        print(line)


def instructions(script, asked):
    """v14's instructions with slide 5's row, the times after it, the deck's name and the figure table brought up to v15."""
    from docx import Document
    doc = Document(str(DECK / "CNS2026_NotAllSpines_v14_6min_INSTRUCTIONS.docx"))
    order = doc.tables[0]
    times = [int(row.cells[2].text.split()[0]) for row in order.rows[1:]]
    times[4] = round(len(script.split()) / WPM * 60) + PAUSE
    clock = 0
    for row, sec in zip(order.rows[1:], times):
        clock += sec
        row.cells[2].paragraphs[0].runs[0].text = f"{sec}\n({clock // 60}:{clock % 60:02d})"
    row = order.rows[5]
    row.cells[1].paragraphs[0].runs[0].text = "Benzel again, with the gold tiers"
    row.cells[3].paragraphs[0].runs[0].text = script
    head = doc.paragraphs[1]
    text = head.text
    text = re.sub(r"public_v14_6min", "public_v15_6min", text)
    text = re.sub(r"Script \d+ words: about \d+:\d+", f"Script {sum(len(r.cells[3].text.split()) for r in order.rows[1:])} words: "
                  f"about {clock // 60}:{clock % 60:02d}", text)
    text = re.sub(r"so \d+ s in hand", f"so {360 - clock} s in hand", text)
    paragraph_text(head, text)
    bullet = doc.paragraphs[7]
    paragraph_text(bullet, bullet.text + " Slide 5 redraws Figs 1.1 and 1.9 on their own axes, with their lines digitised "
                   "from the book's figures and the seventh tick of Fig. 1.1 labelled 1 (L1), under your gold tier lines, "
                   "so the same Thieme permission covers it. " + asked)
    figures = doc.tables[1].rows[4]
    figures.cells[1].paragraphs[0].runs[0].text = "Benzel with the gold tiers, canal rows, angles"
    figures.cells[2].paragraphs[0].runs[0].text = ("Slide 5: Benzel Figs 1.1 and 1.9 redrawn, his lines digitised (Thieme "
                                                    "permission as slide 4), with your gold tier 3 and 4 measurements. "
                                                    "Canal and angles are in arXiv v1")
    doc.save(str(DECK / "CNS2026_NotAllSpines_v15_6min_INSTRUCTIONS.docx"))


def check(path):
    with zipfile.ZipFile(path) as z:
        text = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    text = re.sub(r"<[^>]+>", " ", text)
    for word in FORBIDDEN:
        assert word.lower() not in text.lower(), word


if __name__ == "__main__":
    build()
