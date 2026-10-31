"""paper/cns2026/make_cns_deck_v20_6min.py -- v20 of the 6-minute CNS deck (Greg, 2026-10-31, night).

v19's slide 5 in four clicks, and slide 7 redrawn in Benzel's style. Slide 5 becomes four slides joined by PowerPoint's
Morph transition, so each click moves the lines instead of swapping pictures: the gold tier reweighted to Panjabi's
sex and age in purple; the same line divided into men (blue) and women (red), each matched to his own; our adults of
Tan's cadavers' age and build (green) beside Tan's means; our adults reweighted to Alonge's Nigerian adults (orange)
beside Alonge's lumbar canal. The lines and the reference points are PowerPoint shapes of the same name on every one
of the four slides (prep_v20_benzel_profile.py writes where they fall), which is what Morph needs to move them; a
PowerPoint without Morph (before 2019) falls back to a fade. Built from v14, as v15 to v19 were, so the hand-shortened
notes of slides 2 and 3 stay, and checked against v14's FORBIDDEN words.

    python paper/cns2026/prep_v20_benzel_profile.py
    python paper/cns2026/prep_v20_spinopelvic.py
    python paper/cns2026/make_cns_deck_v20_6min.py
"""
from __future__ import annotations

import copy
import json
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from prep_v20_benzel_profile import PANELS, STATES, references  # noqa: E402

DECK = HERE.parent
SOURCE = DECK / "CNS2026_Schwing_NotAllSpines_public_v14_6min.pptx"
TARGET = DECK / "CNS2026_Schwing_NotAllSpines_public_v20_6min.pptx"
FIG = HERE / "figs_v20_6min"
GEO = json.load(open(FIG / "geometry.json", encoding="utf-8"))
RW = json.load(open(FIG / "reweighted.json", encoding="utf-8"))
ADDED = len(STATES) - 1                     # slides slide 5 gains

# the v14 build's list (make_cns_deck_v14_6min.FORBIDDEN), copied so this runs without v14's imports
FORBIDDEN = ("SchwingNet", "C-arm", "C arm", "carm", "sweep", "proprietary", "fluoro", "render", "DRR", "registration-free",
             "cone-beam", "local shape", "count from the sacrum", "abstain", "abstention", "prove" + "nance", "coupling", "coupled",
             "decoder", "encoder", "stop rule", "chain", "alphabet", "ordinal", "ResNet", "YOLO", "device card",
             "device model", "embedding", "t-SNE", "occlusion", "TotalSegmentator", "BUU", "anchor", "backbone", "neural")
WPM, PAUSE = 140, 3                         # the v14 pace
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
MORPH = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{PML}"><mc:Choice xmlns:p159="{P159}" Requires="p159">'
         '<p:transition spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')


def mean_gap(target, key, series, levels):
    """The mean distance, in mm, between a reweighted line and a reference series over LEVELS."""
    measure = RW["targets"][target]["measures"][key]
    gaps = [abs(measure[lv]["mean"] - series[lv]) for lv in levels if lv in measure and lv in series]
    return sum(gaps) / len(gaps)


def numbers():
    refs = references()
    pj = RW["targets"]
    men, women = pj["panjabi_men"]["measures"]["plate_w_l"]["L3"]["mean"], pj["panjabi_women"]["measures"]["plate_w_l"]["L3"]["mean"]
    canal_w = refs[("panjabi", "canal", "width")]
    alonge_d = refs[("alonge", "canal", "depth")]
    return {
        "n": RW["people"],
        "cw_mean": mean_gap("panjabi", "ct_scw", canal_w, [lv for lv in canal_w if lv != "C2"]),
        "sex_l3": men - women,
        "tan_n": pj["tan"]["who"]["people"], "tan_men": pj["tan"]["who"]["men"],
        "alonge_d": sum(pj["alonge"]["measures"]["ct_scd"][lv]["mean"] - alonge_d[lv] for lv in alonge_d) / len(alonge_d),
    }


def words():
    """Each click's title, caption, references, script and the line to have ready if asked."""
    k = numbers()
    benzel = "Benzel EC. Biomechanics of Spine Stabilization, 3rd ed. Thieme; 2015, ch. 1, Figs 1.1 and 1.9, digitised.   "
    tiers = "Gold tiers 3 and 4: public CT, C2 to the sacrum in view."
    panjabi = "Panjabi MM, et al. Spine 1991;16:861-869 and 888-901; Spine 1992;17:299-306.   "
    return {
        "panjabi": dict(
            title="The same plots, matched to Panjabi's cadavers",
            caption=(f"Black: Benzel, digitised (permission pending). Gray: Panjabi, 12 cadavers. Purple: {k['n']:,} "
                     "adults' whole-spine CTs reweighted to his sex and age. Width at the lower end plate, depth at the upper."),
            refs=benzel + panjabi + tiers,
            script=(f"Same two plots. Black is Benzel, gray is Panjabi's twelve cadavers. Purple is {k['n']:,} adults' "
                    "whole-spine CTs, reweighted to his sex and age: the canal's width sits within a millimeter of his. "
                    "And the textbook has no place for a thirteenth thoracic or a sixth lumbar: here they are."),
            asked=(f"If asked: the canal is read on the CT itself, {k['cw_mean']:.1f} mm from Panjabi's width on average. "
                   "Each of his twelve specimens is matched to our adults of its sex within five years of its age, and "
                   "the line is the mean of the twelve. T13 and L6 are every spine's means, since the series that holds "
                   "the T13s records no sex.")),
        "sexes": dict(
            title="The same line, men and women apart",
            caption=("Blue: our men, matched by age to Panjabi's 8 men. Red: our women, matched to his 4 women. "
                     "Purple was their 8 to 4 blend."),
            refs=benzel + panjabi + tiers,
            script=(f"Divide it by sex. Men run {k['sex_l3']:.0f} millimeters wider than women at L3, and his line "
                    "sits between them."),
            asked=("If asked: the men's and women's lines are the same matching over his men alone and his women "
                   "alone; Panjabi gives no means by sex.")),
        "tan": dict(
            title="Matched to an Asian series",
            caption=(f"Gray: Tan, 10 Chinese Singaporean cadavers, 56 to 77 years. Green: our {k['tan_n']} adults "
                     "inside their age, height and weight."),
            refs="Tan SH, Teo EC, Chua HC. Eur Spine J 2004;13:137-146.   " + benzel + tiers,
            script=("Tan's Chinese Singaporean cadavers, beside our adults of their age and build: smaller bodies, and "
                    "a canal they measured a third shallower than Panjabi's."),
            asked=(f"If asked: Tan reports no sex; our {k['tan_n']} adults inside their ranges of age, height and weight "
                   f"are {k['tan_men']} men. Tan's own comparison puts their canal 20 per cent narrower and 32 per cent "
                   "shallower than Panjabi's, dry bone with a digitiser, so most of that gap is theirs, not the matching.")),
        "alonge": dict(
            title="Matched to an African series",
            caption=("Gray: Alonge, 120 Nigerian adults' CTs, L1 to L5; their width is taken at the widest pedicle cut. "
                     "Orange: our adults reweighted to their sex and age."),
            refs="Alonge OJ, Adeolu AA, Omololu AB, Atalabi OM. Afr J Med Med Sci 2021;50:69-76.   " + benzel + tiers,
            script=("No African American series measures these chords, so here are Alonge's Nigerian adults on CT: "
                    f"their canal depth about {k['alonge_d']:.0f} millimeters under ours."),
            asked=("If asked: of the published series I hold, African American spines appear only in the Hamann-Todd "
                   "skeletons, measured at the pedicles, and in one disc-height study; neither gives these chords. "
                   "Alonge's width is an interpedicular distance at the widest pedicle cut, which is why it climbs to L5; "
                   "Panjabi's is the canal's widest point.")),
    }


def set_text(frame, text):
    """The frame's first run takes the text, so the slide's own type and colour stay; the other runs go."""
    paragraphs = frame.paragraphs
    runs = paragraphs[0].runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)


def paragraph_text(paragraph, text):
    """A Word paragraph's first run takes the text and keeps its style; the other runs go."""
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def clone(prs, source):
    """A new slide on SOURCE's layout holding copies of its text shapes, placed right after it."""
    slide = prs.slides.add_slide(source.slide_layout)
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)
    for shape in source.shapes:
        if shape.has_text_frame:
            tree.append(copy.deepcopy(shape._element))
    ids = prs.slides._sldIdLst
    entry = ids[-1]
    ids.remove(entry)
    ids.insert(next(i for i, e in enumerate(ids) if e.get("id") == str(source.slide_id)) + 1, entry)
    return slide


def plain(shape):
    """No theme style on a drawn shape, so no shadow or theme line comes with it."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def emu(box, point):
    left, top, width, height = box
    return int(left + point[0] * width), int(top + point[1] * height)


def polyline(slide, name, points, colour, width_pt, dash=None):
    builder = slide.shapes.build_freeform(*points[0], scale=1.0)
    builder.add_line_segments(points[1:], close=False)
    shape = builder.convert_to_shape()
    plain(shape)
    shape.name = name
    shape.fill.background()
    shape.line.color.rgb = RGBColor.from_string(colour[1:])
    shape.line.width = Pt(width_pt)
    if dash is not None:
        shape.line.dash_style = dash
    ln = shape.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:round"))
    return shape


def diamond(slide, name, centre, colour, size=Pt(3.4)):
    shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, int(centre[0] - size / 2), int(centre[1] - size / 2), int(size), int(size))
    plain(shape)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(colour[1:])
    shape.line.fill.background()
    return shape


def draw_state(slide, boxes, state, alts):
    """One click: each panel's picture, then its reference points and line, then the gold tier's four lines."""
    for name, box in boxes.items():
        pic = slide.shapes.add_picture(str(FIG / f"{name}_{state}.png"), *box)
        pic.name = f"!!{name}_base"
        pic._element.nvPicPr.cNvPr.set("descr", alts[name])
    for name, box in boxes.items():
        geo = GEO["panels"][name]["states"][state]
        for chord, points in geo["points"].items():
            if len(points) > 1:
                polyline(slide, f"!!{name}_r{chord}", [emu(box, p[1:]) for p in points], GEO["reference"], GEO["ref_pt"],
                         MSO_LINE.ROUND_DOT)
            for level, *xy in points:
                diamond(slide, f"!!{name}_r{chord}_{level}", emu(box, xy), GEO["reference"])
        male, female = geo["colours"]
        for chord in ("w", "d"):
            dash = None if chord == "w" else MSO_LINE.DASH
            for sex, colour in (("m", male), ("f", female)):
                polyline(slide, f"!!{name}_{chord}_{sex}", [emu(box, p) for p in geo["lines"][f"{chord}_{sex}"]], colour,
                         GEO["line_pt"], dash)


def morph(slide):
    """Morph into this slide from the one before, a fade where PowerPoint has no Morph."""
    root = slide._element
    for old in root.findall(qn("p:transition")) + root.findall(f"{{{MC}}}AlternateContent"):
        root.remove(old)
    after = root.find(qn("p:clrMapOvr"))
    if after is None:
        after = root.find(qn("p:cSld"))
    after.addnext(etree.fromstring(MORPH))


def alt_text(state):
    who = {"panjabi": "reweighted to Panjabi's sex and age in purple",
           "sexes": "split into men in blue and women in red, each matched to Panjabi's own",
           "tan": "matched to Tan's cadavers' age and build in green, beside Tan's means in gray",
           "alonge": "reweighted to Alonge's Nigerian adults in orange, beside Alonge's lumbar canal in gray"}[state]
    return {"body": f"Benzel's vertebral body width and depth by level in black, with the gold tier means {who}.",
            "canal": f"Benzel's spinal canal width and depth by level in black, with the gold tier's CT canal means {who}."}


def build():
    texts = words()
    prs = Presentation(str(SOURCE))
    model = [sh for sh in prs.slides[3].shapes if sh.shape_type == 13]
    assert len(model) == 2, len(model)
    boxes = {name: (where.left, where.top, where.width, where.height) for name, where in zip(PANELS, model)}
    first, angles = prs.slides[4], prs.slides[6]
    assert first.shapes[0].text_frame.text.startswith("End-plate width"), first.shapes[0].text_frame.text
    old = first.shapes[1]
    old._element.getparent().remove(old._element)
    spinopelvic(angles)
    slides = [first]
    for _ in range(ADDED):
        slides.append(clone(prs, slides[-1]))
    for slide, state in zip(slides, STATES):
        t = texts[state]
        shapes = [sh for sh in slide.shapes if sh.has_text_frame]
        set_text(shapes[0].text_frame, t["title"])
        set_text(shapes[1].text_frame, t["caption"])
        set_text(shapes[2].text_frame, t["refs"])
        draw_state(slide, boxes, state, alt_text(state))
        notes = slide.notes_slide.notes_text_frame
        notes.text = t["script"]
        notes.add_paragraph().text = t["asked"]
        if slide is not first:
            morph(slide)
    prs.save(str(TARGET))
    check(TARGET)
    instructions(texts)
    print(TARGET.name)
    for state in STATES:
        t = texts[state]
        print(f"[{state}] {t['title']} | {t['caption']} ({len(t['caption'].split())} words) | {t['script']} "
              f"({len(t['script'].split())} words)")


def spinopelvic(slide):
    """Slide 7's picture swapped for the Benzel-style histograms, same place and size."""
    pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
    assert len(pictures) == 1, len(pictures)
    old = pictures[0]
    pic = slide.shapes.add_picture(str(FIG / "spinopelvic_benzel.png"), old.left, old.top, old.width, old.height)
    pic._element.nvPicPr.cNvPr.set("descr", "Histograms of pelvic incidence, sacral slope and pelvic tilt over the "
                                   "gold tier CTs, with the published supine CT and standing film means.")
    old._element.getparent().remove(old._element)
    caption = [sh for sh in slide.shapes if sh.has_text_frame][1]
    caption_text = json.load(open(FIG / "spinopelvic.json", encoding="utf-8"))["caption"]
    set_text(caption.text_frame, caption_text)


def instructions(texts):
    """v14's instructions brought up to v20: slide 5's four clicks in one row, later slides numbered three higher."""
    from docx import Document
    doc = Document(str(DECK / "CNS2026_NotAllSpines_v14_6min_INSTRUCTIONS.docx"))
    order = doc.tables[0]
    script = " [click] ".join(texts[s]["script"] for s in STATES)
    times = [int(row.cells[2].text.split()[0]) for row in order.rows[1:]]
    times[4] = round(len(script.replace("[click]", "").split()) / WPM * 60) + PAUSE * len(STATES)
    clock = 0
    for row, sec in zip(order.rows[1:], times):
        clock += sec
        row.cells[2].paragraphs[0].runs[0].text = f"{sec}\n({clock // 60}:{clock % 60:02d})"
    for row in order.rows[6:]:
        cell = row.cells[0].paragraphs[0].runs[0]
        cell.text = str(int(cell.text) + ADDED)
    row = order.rows[5]
    row.cells[0].paragraphs[0].runs[0].text = f"5-{5 + ADDED}"
    row.cells[1].paragraphs[0].runs[0].text = "Benzel and Panjabi, with the gold tiers: four clicks"
    row.cells[3].paragraphs[0].runs[0].text = script
    head = doc.paragraphs[1]
    text = re.sub(r"public_v14_6min", "public_v20_6min", head.text)
    total = sum(len(r.cells[3].text.replace("[click]", "").split()) for r in order.rows[1:])
    text = re.sub(r"Script \d+ words: about \d+:\d+", f"Script {total} words: about {clock // 60}:{clock % 60:02d}", text)
    text = re.sub(r"so \d+ s in hand", f"so {360 - clock} s in hand", text)
    text = re.sub(r"\d+ slides", f"{len(order.rows) - 1 + ADDED} slides", text)
    paragraph_text(head, text)
    for paragraph in doc.paragraphs[4:6]:
        text = paragraph.text.replace("slide 8", f"slide {8 + ADDED}").replace("Slide 8", f"Slide {8 + ADDED}")
        short = clock - times[7]
        paragraph_text(paragraph, re.sub(r"runs \d+:\d+", f"runs {short // 60}:{short % 60:02d}", text))
    bullet = doc.paragraphs[7]
    paragraph_text(bullet, bullet.text + " Slide 5 redraws Figs 1.1 and 1.9 on their own axes, with their lines digitised "
                   "from the book's figures, under your gold tier lines, so the same Thieme permission covers it. It is "
                   f"now four slides (5 to {5 + ADDED}) joined by Morph: each click moves the lines. Morph needs "
                   "PowerPoint 2019 or Microsoft 365; an older PowerPoint fades between them instead, so check the "
                   "speaker-ready room's copy. " + " ".join(texts[s]["asked"] for s in STATES))
    figures = doc.tables[1]
    for row in figures.rows[1:]:
        cell = row.cells[0].paragraphs[0].runs[0]
        if cell.text == "8":
            cell.text = str(8 + ADDED)
        elif cell.text.startswith("5"):
            cell.text = f"5–{7 + ADDED}"
            row.cells[1].paragraphs[0].runs[0].text = "Benzel with the gold tiers, canal rows, angles"
            row.cells[2].paragraphs[0].runs[0].text = (
                "Slides 5 to 8: Benzel Figs 1.1 and 1.9 redrawn, his lines digitised (Thieme permission as slide 4), "
                "with your gold tier 3 and 4 measurements beside Panjabi, Tan and Alonge's published means. Slide 10's "
                "angle histograms are your own. Canal and angles are in arXiv v1")
    doc.save(str(DECK / "CNS2026_NotAllSpines_v20_6min_INSTRUCTIONS.docx"))


def check(path):
    with zipfile.ZipFile(path) as z:
        text = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    text = re.sub(r"<[^>]+>", " ", text)
    for word in FORBIDDEN:
        assert word.lower() not in text.lower(), word


if __name__ == "__main__":
    build()
